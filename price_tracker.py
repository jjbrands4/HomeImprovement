#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
price_tracker.py - Home wishlist price tracker
==============================================

WHAT IT DOES
  1. Reads the wishlist workbook (Master Sheet + Primary/Secondary Vendor List tabs).
  2. For each selected item it gathers prices in two layers:
       VERIFICATION layer (prices read from the merchant itself)
         - your Master Sheet "Product URLs" (trusted CANDIDATES - validated against GTIN/MPN/title/specs,
           stale redirects detected)                                                  -> verified_direct
         - Best Buy Products API (by SKU, UPC, model or keywords)        [BESTBUY_API_KEY, free] -> verified_direct
         - Shopify stores (auto-detected), retailer site search / sitemaps, DTC vendor sites
           (cached between runs, re-validated each run)                                   -> verified_discovered
         - Playwright browser render, only when a page's data is JavaScript-only (optional)
       DISCOVERY / CORROBORATION layer
         - Google Shopping via SerpApi (optional, cached, skipped when verified pages suffice) -> market_snapshot
         - eBay Browse API (optional - used/refurbished, only if "Open to used" = Yes)
  2b. Master Sheet switches: 'Only Check Primary Links' = Yes -> Product URLs; with >= 3 vendors working no search, with
      fewer a SerpApi general search + a same-vendor search for each vendor whose link failed (vendors with a working
      link are never searched). '!phrase' in Product specifications / Search Keywords = never searched, never matched.
      Run Data records the Search Mode, Primary Link Failed and Expanded Search No Results flags.
  3. Matches identifier-first (GTIN > MPN/SKU > page metadata > title/specs) with hard variant conflicts
     (region, generation, colour, voltage, size, accessory, bundle, sibling model).
  4. Normalises current / regular / shipping / effective price, pack size, condition, availability and
     conditional pricing (coupon, membership, subscription, financing, trade-in ...).
  5. Reconciles offers (stable offer ids, cross-source corroboration, promotion only on independent
     agreement), computes the verified baseline, the reference price (MSRP/regular -> consensus ->
     verified history), market averages, verified-only history stats, the cheapest way to buy the
     Quantity Needed, and deals.
  6. Appends to "Run Data" and "Deals Data" (nothing else in the workbook is modified) and keeps a small
     state folder (tracker_state/) with learned identifiers, discovery cache, offer history.

TRUST RULES
  * Only verified (merchant page / official API), High-confidence, new, ordinary-priced listings feed the
    verified baseline, the 30-day verified median and the historical verified low.
  * Aggregator rows (Google Shopping) never inherit verified status - not even when they stand in for a
    Product URL that could not be fetched.
  * Medium listings count only after independent corroboration promotes them.

USAGE EXAMPLES
  python price_tracker.py --dry-run --force          # run everything, write nothing
  python price_tracker.py --min-priority 4           # only priority 4-5 items
  python price_tracker.py --items 1,3 --force        # specific WishlistItem ids
  python price_tracker.py --run-id my-retry-1        # re-running the same id replaces its rows (idempotent)
  python price_tracker.py --browser off --workers 1  # no Playwright, strictly sequential

All timestamps written to the workbook are UTC.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import time
import uuid
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse

import requests
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pricetrack import pricing
from pricetrack.adapters import (KNOWN_DOMAINS, AdapterContext, BestBuyAdapter, EbayAdapter, EbayClient,
                                 ProductPageAdapter, RetailerDiscoveryAdapter, SerpApiAdapter, SerpApiClient,
                                 ShopifyAdapter, build_query, run_safely)
from pricetrack.adapters.discovery import NO_CRAWL
from pricetrack.fetch import BrowserRenderer, Fetcher
from pricetrack.history import State, history_stats, verified_series, STATE_DIR
from pricetrack.identity import (classify, classify_confidence, gtins_in_text, item_brand, looks_like_sku,  # noqa: F401
                                 normalize_gtin, validate_page)
from pricetrack.models import (MARKET_SNAPSHOT, VERIFIED, VERIFIED_DIRECT, VERIFIED_DISCOVERED, Item,
                               Listing, Outcome, SourceResult, Vendor)
from pricetrack.reconcile import assign_ids, merge_duplicates, promote
from pricetrack.text import (extract_pack_qty, keys_match, norm_text, parse_bulk_sizes, parse_price,  # noqa: F401
                             vendor_key)
from pricetrack.urls import host_of, normalize_url

# =============================================================================
# 1. SETTINGS  (safe to tweak; everything else reads from here)
# =============================================================================

WORKBOOK_GLOB = "Home*Wishlist*.xlsx"

# ---- Deal rules (copied into pricetrack.pricing at start-up) --------------------------------------
DEAL_DISCOUNT = 0.15              # rule A: >= 15% below the reference price (MSRP/regular -> consensus -> history)
TREND_MIN_DISCOUNT = 0.0          # rule B: below the 30-day VERIFIED median by at least this much
TREND_WINDOW_DAYS = 30
TREND_MIN_POINTS = 3              # verified runs needed before the 30-day median is used
MARKET_MIN_SAMPLE = 3
USE_TARGET_RULE = False
REQUIRE_AT_OR_BELOW_TARGET = False
SNAPSHOT_DEALS = True             # Google Shopping rows may still be deals (labelled, ranked after verified ones)
TRUST_LEGACY_BASELINES = False    # count pre-upgrade 'Product URLs' baselines as verified history?

# ---- Noise / sanity filters -----------------------------------------------------------------------
MIN_PRICE_RATIO_OF_TARGET = 0.35
OUTLIER_LOW, OUTLIER_HIGH = 0.4, 2.5
MAX_DEALS_PER_ITEM = 5
ANCHOR_BAND_NEW = (0.60, 1.40)
ANCHOR_BAND_RESALE = (0.30, 1.05)
DEAL_MIN_CONFIDENCE = "High"
OOS_COUNTS_FOR_BASELINE = True

# ---- Cost control -----------------------------------------------------------------------------------
CADENCE_DAYS = {5: 0, 4: 0, 3: 0, 2: 6, 1: 9}
SERP_CACHE_HOURS = 20             # identical SerpApi query inside this window -> cached response (0 credits)
SERP_REDISCOVER_DAYS = 6          # with >= SERP_SKIP_MIN_VERIFIED verified prices, re-query SerpApi only this often
SERP_SKIP_MIN_VERIFIED = 2
SNAPSHOT_VERIFY_MAX = 3           # Google Shopping rows with a direct merchant link that get verified on the merchant page
LINKS_ENOUGH_VENDORS = 3          # 'Only Check Primary Links': this many VENDORS with a working Product URL -> no search at all
LINK_FALLBACK_VERIFY_MAX = 2      # 'Only Check Primary Links': same-vendor Google rows fetched on the merchant page per failed vendor
LINK_FALLBACK_CACHE_DAYS = 30     # a replacement link that worked is tried again first (0 credits) for this long
QUERY_INCLUDE_SPECS = True

# ---- Retrieval --------------------------------------------------------------------------------------
WORKERS = 3                       # conservative parallelism (per-domain limits still apply)
RETAILER_DISCOVERY = True         # look for the product on known retailer sites (cached between runs)

# ---- Workbook layout --------------------------------------------------------------------------------
SHEET_MASTER = "Master Sheet"
SHEET_PRIMARY = "Primary Vendor List"
SHEET_SECONDARY = "Secondary Vendor List"
SHEET_RUN = "Run Data"
SHEET_DEALS = "Deals Data"

RUN_COLS = ["RowDateTime", "runID", "WishlistItem", "Product", "Listings Searched",
            "Unique Websites Searched", "Target Price or better found",
            "Primary Vendor Price or Better", "Deals found", "Primary Vendor Deals",
            "Avg Price (New)", "Avg Price (Resale)", "30d Trend (New)", "30d Trend (Resale)",
            "Lowest Price Found", "Source Notes", "Matching Listings",
            "Avg Sample (New)", "Avg Sample (Resale)",
            "Vendor Baseline (New)", "Baseline Source", "Product URL Status",
            # --- identity / reference / history / quantity / audit ---
            "Reference Price", "Reference Type", "Prior Verified Price", "Change vs Prior",
            "Verified Low", "EWMA (Verified)", "Qty Needed", "Best Qty Plan", "Best Qty Total",
            "Product Identity", "Evidence Mix", "Retrieval Outcomes",
            # --- search-mode flags (also read back from earlier rows to count consecutive runs) ---
            "Search Mode", "Primary Link Failed", "Expanded Search No Results"]
DEALS_COLS = ["runID", "WishlistItem", "Product", "URL", "isPrimaryVendor", "Target Price",
              "Price", "RightProductConfidence",
              "Vendor", "Source", "Condition", "Listed Price", "Pack Qty", "% Below Target",
              "Deal Rule", "Secondary Vendor?", "Secondary Vendor Comments",
              "In Stock Verified", "Listing Title", "Baseline Price", "% Below Baseline",
              # --- provenance ---
              "Evidence", "Match Evidence", "Reference Type", "Regular Price", "Shipping", "Availability",
              "Conditional Pricing", "Corroborated By", "Price Event", "Offer ID", "Retrieved At"]
MONEY_COLS = {"Target Price", "Price", "Listed Price", "Avg Price (New)", "Avg Price (Resale)",
              "30d Trend (New)", "30d Trend (Resale)", "Lowest Price Found", "Vendor Baseline (New)",
              "Baseline Price", "Reference Price", "Prior Verified Price", "Verified Low", "EWMA (Verified)",
              "Best Qty Total", "Regular Price", "Shipping"}
PCT_COLS = {"% Below Target", "% Below Baseline", "Change vs Prior"}
COL_WIDTHS = {"URL": 50, "Source Notes": 60, "Product URL Status": 50, "Listing Title": 55, "Deal Rule": 42,
              "Secondary Vendor Comments": 55, "Vendor": 20, "RowDateTime": 17, "Best Qty Plan": 50,
              "Retrieval Outcomes": 50, "Corroborated By": 40, "Product Identity": 30,
              "Search Mode": 45, "Primary Link Failed": 45, "Expanded Search No Results": 45}


def log(msg: str) -> None:
    print(msg, flush=True)


def _push_settings() -> None:
    """Copy the user-facing settings above into the pricing module (single source of truth here)."""
    for name in ("DEAL_DISCOUNT", "TREND_MIN_DISCOUNT", "MARKET_MIN_SAMPLE", "USE_TARGET_RULE",
                 "REQUIRE_AT_OR_BELOW_TARGET", "SNAPSHOT_DEALS", "MIN_PRICE_RATIO_OF_TARGET", "OUTLIER_LOW",
                 "OUTLIER_HIGH", "MAX_DEALS_PER_ITEM", "ANCHOR_BAND_NEW", "ANCHOR_BAND_RESALE",
                 "DEAL_MIN_CONFIDENCE", "OOS_COUNTS_FOR_BASELINE"):
        setattr(pricing, name, globals()[name])


# =============================================================================
# 2. WORKBOOK READERS
# =============================================================================

def truthy(v) -> bool:
    return str(v).strip().lower() in {"yes", "y", "true", "1", "1.0"}


def split_multi(v) -> list:
    """';' separates the entries of a Master Sheet cell (a line break inside the cell works too)."""
    return [p.strip() for p in re.split(r"[;\uff1b\r\n]+", str(v or "")) if p.strip()]


_BARE_INCH = re.compile(r"(?<=\d)\s*in\b(?![-\s]*\d)(?!-)", re.I)


def tidy_spec(s: str) -> str:
    """'50 in' / '16.4 in' -> '50 inch' (a bare 'in' is never a word in a title, so a spec written that way
    could never be matched). Leaves '4 in 1' and 'in-wall' alone."""
    return _BARE_INCH.sub(" inch", s)


def split_exclusions(parts: list) -> tuple:
    """Entries starting with '!' are EXCLUSIONS ('!Lite' = never use / never match 'Lite'): the phrase is everything
    up to the next ';' or the end of the cell. Returns (positive entries, excluded phrases); repeats that normalise
    to the same words ('50 in', '50 inch', '50"') are collapsed."""
    pos, excl, seen = [], [], set()
    for p in parts:
        neg = p.lstrip().startswith("!")
        text = p.lstrip().lstrip("!").strip()
        if not text:
            continue
        if not looks_like_sku(text):
            text = tidy_spec(text)
        key = (neg, frozenset(norm_text(text).split()) or text.lower())
        if key in seen:
            continue
        seen.add(key)
        (excl if neg else pos).append(text)
    return pos, excl


def norm_header(s) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s or "").lower())


def header_map(ws) -> dict:
    return {norm_header(c.value): c.column for c in ws[1] if c.value not in (None, "")}


def find_col(hm: dict, *names: str) -> Optional[int]:
    for n in names:
        if norm_header(n) in hm:
            return hm[norm_header(n)]
    for n in names:
        for k, v in hm.items():
            if k.startswith(norm_header(n)):
                return v
    return None


_URL_RE = re.compile(r"https?://[^\s;\uff1b<>\"']+", re.I)     # ';' separates URLs; commas inside a URL are kept


def extract_urls(cell) -> list:
    """Every http(s) URL in a cell plus the cell's clickable hyperlink; tracking parameters removed."""
    found = _URL_RE.findall(str(cell.value or ""))
    link = getattr(cell, "hyperlink", None)
    target = getattr(link, "target", None) if link else None
    if target and target.lower().startswith("http"):
        found.append(target)
    return list(dict.fromkeys(normalize_url(u.rstrip(").,")) for u in found))


def load_items(ws) -> list:
    """Master Sheet rows -> Items. Optional identifier columns are used when present:
       'GTIN' / 'UPC' / 'EAN', 'MPN' / 'Model', 'Brand' (none are required)."""
    hm = header_map(ws)
    c = {k: find_col(hm, *v) for k, v in {
        "wid": ("WishlistItem",), "product": ("Product",), "specs": ("Product specifications",),
        "kw": ("Search Keywords",), "loc": ("Location",), "pri": ("Priority",),
        "qty": ("Quantity Needed",), "used": ("Open to used",), "target": ("Target Price",),
        "bulk": ("Is Bulk Option",), "bkw": ("Bulk Keywords",),
        "only": ("Only Check Primary Links", "Only Check Product URLs"),
        "urls": ("Product URLs", "Product URL", "Product Links", "Links"),
        "gtin": ("GTIN", "UPC", "EAN", "GTIN/UPC", "UPC/EAN"), "mpn": ("MPN", "Model Number", "Model", "Manufacturer Part"),
        "brand": ("Brand", "Manufacturer")}.items()}
    if c.get("product") and c.get("mpn") == c.get("product"):
        c["mpn"] = None

    def get(r, k):
        return ws.cell(r, c[k]).value if c.get(k) else None

    items = []
    for r in range(2, ws.max_row + 1):
        product = get(r, "product")
        if not product or not str(product).strip():
            continue
        kws, kw_excl = split_exclusions(split_multi(get(r, "kw")))
        specs, spec_excl = split_exclusions(split_multi(get(r, "specs")))
        urls = extract_urls(ws.cell(r, c["urls"])) if c.get("urls") else []
        only_links = truthy(get(r, "only"))
        if only_links and not urls:
            log(f"  NOTE: item {get(r, 'wid')} has 'Only Check Primary Links' = Yes but no Product URLs - normal search used")
            only_links, links_fallback = False, True
        else:
            links_fallback = False
        skus = [k for k in kws if looks_like_sku(k)]
        spec_kw = [k for k in kws if k not in skus]
        try:
            pri = int(float(get(r, "pri")))
        except (TypeError, ValueError):
            pri = 3
        gtins = {g for g in (normalize_gtin(x) for x in split_multi(get(r, "gtin"))) if g}
        gtins |= {g for k in kws for g in gtins_in_text(k)}
        items.append(Item(
            wid=get(r, "wid"), product=str(product).strip(), specs=specs, keywords=kws,
            location=str(get(r, "loc") or "").strip(), priority=pri, qty=get(r, "qty"),
            open_used=truthy(get(r, "used")), target=parse_price(get(r, "target")),
            bulk=truthy(get(r, "bulk")), bulk_sizes=parse_bulk_sizes(get(r, "bkw")),
            skus=skus, spec_phrases=specs + spec_kw,
            urls=urls, gtins=gtins, mpns=split_multi(get(r, "mpn")), brand=str(get(r, "brand") or "").strip(),
            exclude=list(dict.fromkeys(spec_excl + kw_excl)), only_links=only_links,
            links_only_fallback=links_fallback))
    return items


def load_primary_vendors(ws) -> list:
    hm = header_map(ws)
    cv, cd, cdom = find_col(hm, "Vendor"), find_col(hm, "isDirectToConsumer"), find_col(hm, "Domain")
    out = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, cv).value if cv else None
        if not name:
            continue
        key = vendor_key(str(name))
        domain = str(ws.cell(r, cdom).value or "").strip() if cdom else ""
        domain = domain or next((d for k, d in KNOWN_DOMAINS.items() if keys_match(key, k)), "")
        out.append(Vendor(str(name).strip(), key, truthy(ws.cell(r, cd).value) if cd else False, domain))
    return out


def load_secondary_keys(ws) -> list:
    hm = header_map(ws)
    cv = find_col(hm, "Vendor")
    return [vendor_key(str(ws.cell(r, cv).value)) for r in range(2, ws.max_row + 1)
            if cv and ws.cell(r, cv).value]


def _to_dt(v) -> Optional[datetime]:
    if isinstance(v, datetime):
        return v
    try:
        return datetime.fromisoformat(str(v))
    except (TypeError, ValueError):
        return None


def load_history(ws) -> list:
    """Past Run Data rows (used for verified history, the resale trend and cadence)."""
    hm = header_map(ws)
    col = {k: find_col(hm, n) for k, n in {
        "dt": "RowDateTime", "run": "runID", "wid": "WishlistItem", "avg_new": "Avg Price (New)",
        "avg_res": "Avg Price (Resale)", "n_new": "Avg Sample (New)", "n_res": "Avg Sample (Resale)",
        "baseline": "Vendor Baseline (New)", "baseline_src": "Baseline Source",
        "link_failed": "Primary Link Failed", "no_results": "Expanded Search No Results"}.items()}
    out = []
    for r in range(2, ws.max_row + 1):
        dt = _to_dt(ws.cell(r, col["dt"]).value) if col["dt"] else None
        if dt is None:
            continue
        g = lambda k: ws.cell(r, col[k]).value if col[k] else None   # noqa: E731
        out.append({"wid": str(g("wid")).strip(), "dt": dt, "run": str(g("run") or ""),
                    "avg_new": parse_price(g("avg_new")), "avg_res": parse_price(g("avg_res")),
                    "n_new": parse_price(g("n_new")), "n_res": parse_price(g("n_res")),
                    "baseline": parse_price(g("baseline")), "baseline_src": str(g("baseline_src") or ""),
                    "link_failed": str(g("link_failed") or ""), "no_results": str(g("no_results") or "")})
    return out


def flag_streak(history: list, wid, field: str, run_id: str = "") -> int:
    """How many of this item's most recent earlier runs (newest first) already carried a 'Yes...' in Run Data column `field`."""
    rows = sorted((h for h in history if h["wid"] == str(wid).strip() and h.get("run") != run_id),
                  key=lambda h: h["dt"], reverse=True)
    n = 0
    for h in rows:
        if str(h.get(field) or "").startswith("Yes"):
            n += 1
        else:
            break
    return n


def flag_text(items: list, history: list, wid, field: str, run_id: str) -> str:
    """'No' or 'Yes: a; b' - with '(run N in a row)' when earlier Run Data rows were flagged too."""
    if not items:
        return "No"
    n = flag_streak(history, wid, field, run_id)
    return "Yes" + (f" (run {n + 1} in a row)" if n else "") + ": " + "; ".join(items)


def resale_trend(history: list, wid, now: datetime, exclude_run: str = "") -> Optional[float]:
    """Mean of stored resale averages (runs with >= MARKET_MIN_SAMPLE listings) over the trend window."""
    cutoff = now - timedelta(days=TREND_WINDOW_DAYS)
    vals = [h["avg_res"] for h in history if h["wid"] == str(wid).strip() and h["dt"] >= cutoff and h["avg_res"]
            and h.get("run") != exclude_run and (h.get("n_res") is None or h["n_res"] >= MARKET_MIN_SAMPLE)]
    return round(sum(vals) / len(vals), 2) if len(vals) >= TREND_MIN_POINTS else None


# =============================================================================
# 3. RETRIEVAL ORCHESTRATION
# =============================================================================

@dataclass
class Context:
    run_id: str
    now: datetime
    primary: list
    secondary_keys: list
    history: list
    state: State
    fetcher: Fetcher
    adapters: dict
    actx: AdapterContext
    args: argparse.Namespace


def url_vendor_name(host: str, primary: list) -> str:
    """'bestbuy.com' -> 'Best Buy' (your Primary Vendor name) so isPrimaryVendor is set correctly."""
    for v in primary:
        if v.domain and (host == host_of("https://" + v.domain) or host.endswith("." + host_of("https://" + v.domain))):
            return v.name
    hkey = vendor_key(host.split(":")[0])
    return next((v.name for v in primary if keys_match(hkey, v.key)), host)


def vendor_name_matches_product(vendor_name: str, item: Item) -> bool:
    """'Philips Hue Direct' ~ 'Hue Color Slim Downlight' (only then is a DTC site worth checking)."""
    stop = {"the", "inc", "llc", "store", "shop", "direct", "official", "online"}
    toks = [t for t in re.split(r"[^a-z0-9]+", vendor_name.lower()) if len(t) >= 3 and t not in stop]
    hay = norm_text(item.product + " " + " ".join(item.keywords) + " " + item.brand)
    hay_set, hay_compact = set(hay.split()), hay.replace(" ", "")
    return any(t in hay_set or (len(t) >= 6 and t in hay_compact) for t in toks)


def _parallel(ctx: Context, jobs: list) -> list:
    """jobs: [(source, target, fn, args, kwargs)] -> SourceResults (isolated, conservative parallelism)."""
    if not jobs:
        return []
    workers = max(1, min(ctx.args.workers, len(jobs)))
    if workers == 1:
        return [run_safely(fn, s, t, *a, **kw) for s, t, fn, a, kw in jobs]
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="src") as ex:
        futs = [ex.submit(run_safely, fn, s, t, *a, **kw) for s, t, fn, a, kw in jobs]
        return [f.result() for f in futs]


def product_url_phase(item: Item, ctx: Context) -> tuple:
    """Price the Master Sheet Product URLs (trusted candidates) and learn identifiers from those that validate."""
    page = ctx.adapters["page"]
    jobs = []
    for u in item.urls:
        host = host_of(u)
        vname = url_vendor_name(host, ctx.primary)
        jobs.append(("page", host, page.fetch, (u, item, ctx.actx), {"vendor": vname, "evidence": VERIFIED_DIRECT}))
    results = _parallel(ctx, jobs)
    listings, status, failed, learned = [], {}, [], []
    for u, res in zip(item.urls, results):
        host = host_of(u)
        url_gtins = gtins_in_text(urlparse(u).path)
        for l in res.listings:
            l.from_url = True
            if not l.gtins and url_gtins and len(res.listings) == 1:
                l.gtins = set(url_gtins)                   # e.g. philips-hue.com/.../046677609405
            l.pack_qty = max(1, int(l.pack_qty_hint or extract_pack_qty(l.title) or 1))
        # validate this page's identity now (identifier learning feeds the discovery phase)
        best, verdicts = None, []
        for l in res.listings:
            m = validate_page(item, l.title, l.gtins, l.mpns, brand_hint=l.brand, domain_brand=f"{l.vendor} {host}",
                              slug=l.page_slug, pack_qty=l.pack_qty, color=l.color)
            verdicts.append(m)
            if m.confidence == "High" and (best is None or l.pack_qty < best[0].pack_qty):
                best = (l, m)
        if res.listings and not best and all(v.confidence == "Low" for v in verdicts):
            res = SourceResult(res.source, res.target, Outcome.IDENTITY_MISMATCH,
                               f"page is not this product - {verdicts[0].reason}", res.listings)
            results[item.urls.index(u)] = res
        if best and best[0].pack_qty == 1:
            l, m = best
            learned += ctx.state.learn(item, l.gtins, l.mpns, l.brand, normalize_url(u))
        elif best:
            learned += ctx.state.learn(item, (), best[0].mpns, best[0].brand, normalize_url(u))
        listings += res.listings
        if res.listings:
            first = sorted(res.listings, key=lambda x: x.pack_qty)[0]
            stock = {True: "in stock", False: "OUT OF STOCK", None: "stock unknown"}[first.in_stock]
            ident = (f"identity validated ({best[1].evidence or 'title'})" if best else
                     "IDENTITY MISMATCH (excluded)" if res.outcome == Outcome.IDENTITY_MISMATCH else
                     "identity unconfirmed - Medium, needs corroboration (not trusted)")
            status[u] = (f"{host}: [{res.outcome}] ${first.price:,.2f} {stock} - {res.detail}; {ident}"
                         + (f" (+{len(res.listings) - 1} variants)" if len(res.listings) > 1 else ""))
            if res.outcome == Outcome.IDENTITY_MISMATCH:       # the link now shows a different product: not a usable link
                failed.append((u, host, url_vendor_name(host, ctx.primary), res.outcome))
        else:
            status[u] = f"{host}: [{res.outcome}] FAILED - {res.detail}"
            failed.append((u, host, url_vendor_name(host, ctx.primary), res.outcome))
    return listings, status, failed, learned, results


def verification_phase(item: Item, ctx: Context, covered: set) -> list:
    """Retailer adapters (identifier-first): Best Buy API, DTC + known-retailer discovery (cached)."""
    a, jobs = ctx.args, []
    bb = ctx.adapters["bestbuy"]
    if "bestbuy.com" not in covered and bb.available()[0] and any(keys_match(v.key, "bestbuy") for v in ctx.primary):
        bbv = next(v.name for v in ctx.primary if keys_match(v.key, "bestbuy"))
        jobs.append(("bestbuy_api", "search", bb.search, (item, ctx.actx), {"vendor": bbv}))
    disc = ctx.adapters["discovery"]
    for v in ctx.primary:
        dom = host_of("https://" + v.domain) if v.domain else ""
        if not dom or dom in covered or dom in NO_CRAWL:
            continue
        if v.is_dtc:
            if vendor_name_matches_product(v.name, item):
                jobs.append(("discovered", v.name, disc.search, (item, ctx.actx), {"vendor": v, "allow_engines": True}))
        elif RETAILER_DISCOVERY and not a.no_discovery:
            jobs.append(("discovered", v.name, disc.search, (item, ctx.actx), {"vendor": v, "allow_engines": False}))
    return _parallel(ctx, jobs)


def verify_snapshots(item: Item, ctx: Context, snaps: list, covered: set) -> list:
    """Aggregators are discovery: when a Google Shopping row links straight to a merchant, fetch that
    merchant page and price it there (verified_discovered); remember the URL for next runs."""
    out, done = [], set()
    cands = [s for s in snaps if s.url.startswith("http") and "google." not in host_of(s.url)
             and host_of(s.url) not in covered and host_of(s.url) not in NO_CRAWL
             and classify(item, s.title).confidence in ("High", "Medium")]
    jobs = []
    for s in cands:
        dom = host_of(s.url)
        if dom in done or len(jobs) >= SNAPSHOT_VERIFY_MAX:
            continue
        done.add(dom)
        jobs.append(("page", dom, ctx.adapters["page"].fetch, (s.url, item, ctx.actx),
                     {"vendor": s.vendor or dom, "evidence": VERIFIED_DISCOVERED}))
    for (src, dom, *_), res in zip(jobs, _parallel(ctx, jobs)):
        good = [l for l in res.listings if classify(item, l.title, gtins=l.gtins, mpns=l.mpns,
                                                    pack_qty=l.pack_qty_hint or 1).confidence in ("High", "Medium")]
        for l in good:
            l.source = "discovered"
        if good:
            ctx.state.discovery_put(item, dom, url=good[0].url, ok=True)
        out.append(SourceResult("snapshot-verify", dom, res.outcome if good or not res.listings else Outcome.IDENTITY_MISMATCH,
                                res.detail, good))
    return out


def fill_failed_urls(item: Item, failed: list, snaps: list, status: dict) -> None:
    """A Product URL that could not be fetched may be covered by the SAME merchant's Google Shopping row -
    as a market_snapshot (coverage only; it never becomes the verified baseline)."""
    for url, host, vname, outcome in failed:
        keys = {vendor_key(vname), vendor_key(host)}
        cands = [s for s in snaps if any(keys_match(vendor_key(s.vendor), k) for k in keys) and s.condition == "new"
                 and classify(item, s.title).confidence in ("High", "Medium")]
        if not cands:
            continue
        best = min(cands, key=lambda s: (classify(item, s.title).confidence != "High", extract_pack_qty(s.title)))
        best.vendor = vname
        best.seller_comment = (f"Google Shopping row for your Product URL ({host} {outcome}); "
                               f"market snapshot, price not verified on the merchant page")
        status[url] = status.get(url, "") + f" -> ${best.price:,.2f} via Google Shopping fallback (market_snapshot, not verified)"


def link_fallback_phase(item: Item, failed: list, ctx: Context, url_status: dict, ok_hosts: set) -> tuple:
    """'Only Check Primary Links' = Yes. For each VENDOR whose Product URL(s) could not be used, look for that same
    vendor's own listing of the product through SerpApi (Google Shopping) - e.g. a dead Best Buy link -> the Best Buy
    row Google lists for the exact product. Nothing else is searched.
      1. a replacement link that worked on an earlier run is re-fetched first (0 credits)
      2. SerpApi rows from that vendor only (cached 20 h), identity-checked; their direct merchant link is fetched and
         priced on the merchant page itself (verified_discovered)
      3. no fetchable link -> the vendor's Google row stays a market_snapshot, clearly labelled as not verified
    A vendor with another Product URL that did price is not searched (no credit spent).
    Returns (listings, results, notes, empty) - `empty` names every vendor whose search yielded nothing usable."""
    a = ctx.args
    serp, page = ctx.adapters["serpapi"], ctx.adapters["page"]
    ok_keys = {vendor_key(h) for h in ok_hosts}
    out, results, notes, tried, empty = [], [], [], set(), []
    packs = pricing.allowed_packs(item)

    def verified(res) -> list:
        return [x for x in res.listings
                if classify(item, x.title, gtins=x.gtins, mpns=x.mpns, pack_qty=x.pack_qty_hint or 1).confidence
                in ("High", "Medium")]

    for url, host, vname, outcome in failed:
        key = vendor_key(host)
        if key in tried or key in ok_keys:
            continue
        tried.add(key)
        label = vname or host
        # 1. cached replacement link
        cached = ctx.state.discovery_get(item, host)
        if (cached and cached.get("ok") and cached.get("url") and normalize_url(cached["url"]) != normalize_url(url)
                and ctx.state.age_days(cached.get("ts")) <= LINK_FALLBACK_CACHE_DAYS):
            res = run_safely(page.fetch, "page", host, cached["url"], item, ctx.actx, vendor=label, evidence=VERIFIED_DISCOVERED)
            good = verified(res)
            if good:
                for x in good:
                    x.source, x.method = "discovered", f"{x.method} (replacement link from an earlier run)"
                out += good
                results.append(SourceResult("link-fallback", label, Outcome.SUCCESS, "cached replacement link", good))
                url_status[url] = (url_status.get(url, f"{host}: [{outcome}] FAILED") +
                                   f" -> replacement link re-verified: {good[0].url} ${good[0].price:,.2f}")
                notes.append(f"{label}: Product URL failed ({outcome}); using replacement link {good[0].url}")
                continue
        # 2. SerpApi, this vendor only
        if a.no_serpapi:
            url_status[url] = url_status.get(url, f"{host}: [{outcome}]") + " -> same-vendor fallback skipped (--no-serpapi)"
            results.append(SourceResult("link-fallback", label, Outcome.SKIPPED, "--no-serpapi"))
            continue
        r = run_safely(serp.search_vendor, "serpapi", label, item, ctx.actx, vendor=label, host=host,
                       cache_hours=SERP_CACHE_HOURS)
        results.append(r)
        if r.outcome not in (Outcome.SUCCESS,):
            why = r.detail or r.outcome
            url_status[url] = url_status.get(url, f"{host}: [{outcome}]") + f" -> same-vendor fallback: {r.outcome} ({why})"
            notes.append(f"{label}: Product URL failed ({outcome}); same-vendor fallback {r.outcome} ({why})"[:200])
            if r.outcome == Outcome.NO_MATCH:
                empty.append(f"{label} (same-vendor Google Shopping)")
            continue
        for l in r.listings:
            pricing.normalize_prices(l)
        if "." in label and r.listings:
            label = r.listings[0].vendor or label                    # 'etsy.com' -> 'Etsy' (the name Google shows)
        cands = [l for l in r.listings
                 if (l.condition == "new" or item.open_used) and (packs is None or l.pack_qty in packs)
                 and classify(item, l.title, gtins=l.gtins, mpns=l.mpns, pack_qty=l.pack_qty).confidence != "Low"]
        direct = lambda l: l.url.startswith("http") and "google." not in host_of(l.url)       # noqa: E731
        cands.sort(key=lambda l: (classify(item, l.title, gtins=l.gtins, mpns=l.mpns, pack_qty=l.pack_qty).confidence != "High",
                                  l.pack_qty != 1, not direct(l), l.unit_price))
        if not cands:
            url_status[url] = url_status.get(url, f"{host}: [{outcome}]") + f" -> no matching {label} listing found via Google Shopping"
            notes.append(f"{label}: Product URL failed ({outcome}); no matching {label} listing on Google Shopping")
            empty.append(f"{label} (same-vendor Google Shopping)")
            continue
        fetched = []
        for l in [c for c in cands if direct(c)][:LINK_FALLBACK_VERIFY_MAX]:
            res = run_safely(page.fetch, "page", host_of(l.url), l.url, item, ctx.actx, vendor=label, evidence=VERIFIED_DISCOVERED)
            good = verified(res)
            if good:
                for x in good:
                    x.source, x.method = "discovered", f"{x.method} (SerpApi same-vendor fallback)"
                fetched = good
                ctx.state.discovery_put(item, host, url=normalize_url(good[0].url), ok=True)
                break
        if fetched:
            out += fetched
            url_status[url] = (url_status.get(url, f"{host}: [{outcome}] FAILED") +
                               f" -> replacement {label} listing verified on the merchant page: {fetched[0].url} ${fetched[0].price:,.2f}")
            notes.append(f"{label}: Product URL failed ({outcome}); replacement {fetched[0].url} (update the Master Sheet link)")
            continue
        best = {}                                                       # cheapest-confidence row per pack size
        for l in cands:
            best.setdefault(l.pack_qty, l)
        for l in best.values():
            l.vendor = label
            l.seller_comment = (f"Google Shopping row from {label} standing in for your Product URL ({host} {outcome}); "
                                f"market snapshot, price not verified on the merchant page")
            out.append(l)
        first = next(iter(best.values()))
        url_status[url] = (url_status.get(url, f"{host}: [{outcome}] FAILED") +
                           f" -> ${first.unit_price:,.2f} from {label} via Google Shopping (market_snapshot, not verified"
                           + (f"; link {first.url}" if direct(first) else "") + ")")
        notes.append(f"{label}: Product URL failed ({outcome}); {label} Google row used as market_snapshot")
    return out, results, notes, empty


def serp_general_phase(item: Item, ctx: Context, listings: list, failed: list, url_status: dict, covered: set,
                       notes: list, results: list, empty: list, links_only: bool = False,
                       skip_keys: frozenset = frozenset()) -> list:
    """General Google Shopping search (SerpApi) for the product + merchant verification of rows that link straight to a
    merchant. Returns the market-snapshot rows. links_only: always run (the SerpApi policy below is for normal items) and
    drop rows from vendors (`skip_keys`) that already have a working Product URL - their price is known from the page."""
    a, serp = ctx.args, ctx.adapters["serpapi"]
    if a.no_serpapi:
        notes.append("SerpApi: skipped (--no-serpapi)")
        return []
    if not serp.available()[0] and serp.client.reason == "SERPAPI_KEY not set":
        notes.append(f"SerpApi: {Outcome.API} ({serp.client.reason})")
        results.append(SourceResult("serpapi", "Google Shopping", Outcome.API, serp.client.reason))
        return []
    n_verified = sum(1 for l in listings if l.evidence in VERIFIED and l.in_stock is not False)
    # Policy: when verified merchant pages already price the item and discovery ran recently, don't
    # spend a credit - but a cached response (e.g. a re-run of the same day) is free and still used.
    skip = (not links_only and not a.force_discovery and n_verified >= SERP_SKIP_MIN_VERIFIED and not failed
            and ctx.state.serp_age_days(item) < SERP_REDISCOVER_DAYS)
    r = run_safely(serp.search, "serpapi", "Google Shopping", item, ctx.actx, include_specs=QUERY_INCLUDE_SPECS,
                   cache_hours=SERP_CACHE_HOURS, cache_only=skip)
    if skip and r.outcome == Outcome.SKIPPED:
        r.detail = (f"{n_verified} verified merchant prices; last discovery {ctx.state.serp_age_days(item):.1f}d ago "
                    f"(< {SERP_REDISCOVER_DAYS}d) - credit saved")
    results.append(r)
    snaps = r.listings
    if skip_keys:
        kept = [s for s in snaps if not any(keys_match(vendor_key(s.vendor), k) for k in skip_keys)]
        if len(kept) != len(snaps):
            notes.append(f"{len(snaps) - len(kept)} Google row(s) from vendors with a working Product URL ignored")
        snaps = kept
    notes.append(f"SerpApi: {r.outcome} {len(snaps)} listings" + (f" ({r.detail})" if r.detail else ""))
    if r.outcome == Outcome.NO_MATCH or (r.outcome == Outcome.SUCCESS and not snaps):
        empty.append("Google Shopping (general search)")
    if snaps and not a.no_direct:
        vr = verify_snapshots(item, ctx, snaps, covered | {host_of(l.url) for l in listings if l.evidence in VERIFIED})
        results += vr
        for x in vr:
            listings += x.listings
        if any(x.listings for x in vr):
            notes.append("verified on merchant page: " + ", ".join(x.target for x in vr if x.listings))
    if failed and snaps and not links_only:
        fill_failed_urls(item, failed, snaps, url_status)
    return snaps


def _outcome_summary(results: list) -> str:
    cnt = Counter(r.outcome for r in results)
    parts = []
    for oc in Outcome.ALL:
        if not cnt.get(oc):
            continue
        who = sorted({r.target for r in results if r.outcome == oc})
        parts.append(f"{oc} {cnt[oc]}" + ("" if oc == Outcome.SUCCESS else f" ({', '.join(who)[:80]})"))
    return " | ".join(parts)


def process_item(item: Item, ctx: Context) -> tuple:
    """Run every layer for one item. Returns (run_row, deal_rows, deals)."""
    a = ctx.args
    ctx.state.apply_learned(item)
    notes, results = [], []
    links_only = item.only_links          # Master Sheet 'Only Check Primary Links' = Yes

    # --- 1. Product URLs (trusted candidates) --------------------------------------------------------
    listings, url_status, failed, learned = [], {}, [], []
    if item.urls and not a.no_direct:
        ls, url_status, failed, learned, res = product_url_phase(item, ctx)
        listings += ls
        results += res
    elif links_only:                      # --no-direct: no merchant pages, so every vendor goes to the SerpApi fallback
        failed = [(u, host_of(u), url_vendor_name(host_of(u), ctx.primary), "skipped") for u in item.urls]
    covered = {host_of(u) for u in item.urls} if not a.no_direct else set()
    if learned:
        notes.append("learned " + ", ".join(learned))

    # --- 2. Verification adapters (identifier-first, cached discovery) -------------------------------
    if links_only:
        notes.append("Only Check Primary Links = Yes: vendor lists, retailer discovery, Best Buy keyword search and eBay skipped")
        results.append(SourceResult("search", "vendor lists + retailer discovery + eBay", Outcome.SKIPPED,
                                    "Only Check Primary Links = Yes"))
    elif not a.no_direct:
        for r in verification_phase(item, ctx, covered):
            results.append(r)
            listings += r.listings
            if r.outcome in (Outcome.SUCCESS, Outcome.UNAVAILABLE, Outcome.IDENTITY_MISMATCH, Outcome.PARSER):
                notes.append(r.note()[:160])          # (no_match / skipped / blocked are summarised in Retrieval Outcomes)

    # --- 3. Discovery / corroboration: SerpApi (optional, cached, policy-gated) ----------------------
    snaps, empty = [], []
    failed_urls = {f[0] for f in failed}
    ok_hosts = {host_of(u) for u in item.urls if u not in failed_urls}          # vendors with >= 1 link that priced
    working = {vendor_key(h) for h in ok_hosts if vendor_key(h)}
    if links_only and len(working) >= LINKS_ENOUGH_VENDORS:
        mode = f"Product URLs only ({len(working)} vendor links work, >= {LINKS_ENOUGH_VENDORS}: no search)"
        notes.append(f"Only Check Primary Links: {len(working)} vendors with a working link - no search run")
        results.append(SourceResult("search", "expanded search", Outcome.SKIPPED, mode))
    elif links_only:
        # < 3 vendors with a working link: SerpApi general search + a same-vendor search for every vendor whose
        # link(s) all failed. A vendor that has a working link is never searched.
        mode = f"Product URLs + expanded SerpApi search ({len(working)} of {LINKS_ENOUGH_VENDORS} vendor links work)"
        fb_ls, fb_res, fb_notes, fb_empty = link_fallback_phase(item, failed, ctx, url_status, ok_hosts)
        listings += fb_ls
        results += fb_res
        notes += fb_notes
        empty += fb_empty
        snaps = serp_general_phase(item, ctx, listings, failed, url_status, covered, notes, results, empty,
                                   links_only=True, skip_keys=frozenset(working))
    else:
        mode = ("Normal search (Only Check Primary Links = Yes but no Product URLs are listed)"
                if item.links_only_fallback else "Normal search")
        if item.links_only_fallback:
            notes.append("Only Check Primary Links = Yes but no Product URLs listed: normal search used")
        snaps = serp_general_phase(item, ctx, listings, failed, url_status, covered, notes, results, empty)

    # --- 4. eBay (optional; only when open to used / wider search) -----------------------------------
    ebay_ls = []
    if item.open_used and not links_only:
        eb = ctx.adapters["ebay"]
        if a.no_ebay:
            notes.append("eBay: skipped (--no-ebay)")
        else:
            r = run_safely(eb.search, "ebay", "eBay", item, ctx.actx, query=build_query(item, include_specs=False))
            results.append(r)
            ebay_ls = r.listings
            notes.append(f"eBay: {r.outcome}" + (f" {len(ebay_ls)} listings" if ebay_ls else f" ({r.detail})"))
            if r.outcome == Outcome.SUCCESS:   # the API copy carries seller ratings: drop Google's eBay rows
                snaps = [s for s in snaps if vendor_key(s.vendor) != "ebay"]
    listings = listings + snaps + ebay_ls

    # --- 5. Score -> reconcile -> baselines -----------------------------------------------------------
    dtc_keys = {v.key for v in ctx.primary if v.is_dtc} | {vendor_key(item_brand(item))}
    pricing.score_listings(item, listings, ctx.primary, ctx.secondary_keys)
    assign_ids(item, listings)
    n_before = len(listings)
    listings = merge_duplicates(listings)
    if n_before != len(listings):
        notes.append(f"{n_before - len(listings)} duplicate row(s) merged as corroboration "
                     f"(same offer from another source, e.g. a Google row for a store already priced directly)")
    promoted = promote(item, listings)
    if promoted:
        notes.append(f"{promoted} Medium listing(s) promoted on independent corroboration")

    baseline, baseline_src, pack_baselines = pricing.verified_baseline(listings)
    series = verified_series(ctx.history, item.wid, exclude_run=ctx.run_id, trust_legacy=TRUST_LEGACY_BASELINES)
    hist = history_stats(series, ctx.state.trusted_units(item, exclude_run=ctx.run_id), ctx.now,
                         TREND_WINDOW_DAYS, TREND_MIN_POINTS)
    ref = pricing.reference_price(item, listings, hist, dtc_keys)
    pricing.apply_band(item, listings, baseline or ref[0])
    pricing.mark_trusted(listings)
    baseline, baseline_src, pack_baselines = pricing.verified_baseline(listings)   # after the sanity band
    stats = pricing.market_stats(listings)
    res_trend = resale_trend(ctx.history, item.wid, ctx.now, ctx.run_id)
    deals = pricing.find_deals(item, listings, stats, hist, ref, baseline, pack_baselines, res_trend)
    plan = pricing.quantity_plan(item, listings)

    # --- 6. History: offer change events + observations (idempotent per run id) ---------------------
    for l in listings:
        if l.confidence in ("High", "Medium") and (l.eligible or l.cond_ok or l.from_url):
            l.price_event = ctx.state.offer_event(l, item, ctx.run_id)
    ctx.state.record_observations(ctx.run_id, item, listings)

    # --- 7. Output rows ------------------------------------------------------------------------------
    rep = [l for l in listings if l.reportable and l.confidence == DEAL_MIN_CONFIDENCE]
    tgt = item.target
    at_target = [l for l in rep if tgt and l.unit_price <= tgt]
    lowest = min((l.unit_price for l in rep), default=None)
    if baseline and any(l.trusted and l.in_stock is False for l in listings) and "sold out" in baseline_src:
        notes.append("Verified baseline uses a sold-out single-unit list price")
    for pool, label in (("new", "New"), ("res", "Resale")):
        s = stats[pool]
        if s and s["n"] < MARKET_MIN_SAMPLE and not (pool == "new" and baseline):
            notes.append(f"Avg ({label}) from only {s['n']} listing(s): recorded, but deal rules need {MARKET_MIN_SAMPLE}+")
    if item.urls and a.no_direct and not links_only:
        url_status = {u: "skipped (--no-direct)" for u in item.urls}

    change = None
    if baseline and hist.get("prior"):
        change = round((baseline - hist["prior"]) / hist["prior"], 4)
        if abs(change) >= 0.02:
            notes.append(f"verified price {'down' if change < 0 else 'up'} {abs(change):.1%} vs prior "
                         f"${hist['prior']:.2f} ({hist['prior_dt']:%Y-%m-%d})")
    ev = Counter(l.evidence for l in listings if l.eligible)
    evidence_mix = " | ".join(f"{k} {ev[k]}" for k in (VERIFIED_DIRECT, VERIFIED_DISCOVERED, MARKET_SNAPSHOT) if ev.get(k))
    cond = [l for l in listings if l.cond_ok]
    if cond:
        notes.append(f"{len(cond)} conditional price(s) kept out of averages ({', '.join(sorted({l.conditional for l in cond}))})")
    identity = []
    if item.all_gtins:
        identity.append("GTIN " + ", ".join(sorted(g.lstrip("0") for g in item.all_gtins)))
    if item.all_mpns:
        identity.append("MPN/SKU " + ", ".join(item.all_mpns[:3]))
    if item.brand or item.learned_brand:
        identity.append("brand " + (item.brand or item.learned_brand))

    # Run Data flags (also read back next run to count consecutive runs): failed primary link / empty expanded search
    if not item.urls:
        link_flag = "No links listed"
    elif a.no_direct:
        link_flag = "Not checked (--no-direct)"
    else:
        link_flag = flag_text(list(dict.fromkeys(f"{h} [{oc}]" for _, h, _, oc in failed)), ctx.history, item.wid,
                              "link_failed", ctx.run_id)
    searched = any(r.source in ("serpapi", "link-fallback") and r.outcome in (Outcome.SUCCESS, Outcome.NO_MATCH, Outcome.NETWORK)
                   for r in results)
    if searched:
        none_flag = flag_text(empty, ctx.history, item.wid, "no_results", ctx.run_id)
    else:
        why = ("SerpApi unavailable" if any(r.outcome == Outcome.API for r in results) else
               "--no-serpapi" if a.no_serpapi else "not needed" if links_only else "")
        none_flag = "Not run" + (f" ({why})" if why else "")

    run_row = {
        "RowDateTime": ctx.now, "runID": ctx.run_id, "WishlistItem": item.wid, "Product": item.product,
        "Listings Searched": len(listings), "Matching Listings": sum(1 for l in listings if l.eligible),
        "Unique Websites Searched": len({l.vkey for l in listings if l.vkey}),
        "Target Price or better found": len(at_target),
        "Primary Vendor Price or Better": sum(1 for l in at_target if l.is_primary),
        "Deals found": len(deals), "Primary Vendor Deals": sum(1 for l in deals if l.is_primary),
        "Avg Price (New)": stats["new"]["avg"] if stats["new"] else None,
        "Avg Price (Resale)": stats["res"]["avg"] if stats["res"] else None,
        "Avg Sample (New)": stats["new"]["n"] if stats["new"] else None,
        "Avg Sample (Resale)": stats["res"]["n"] if stats["res"] else None,
        "30d Trend (New)": hist.get("median"), "30d Trend (Resale)": res_trend,
        "Vendor Baseline (New)": baseline, "Baseline Source": baseline_src,
        "Product URL Status": " | ".join(url_status.values())[:900] if url_status else
                              ("none listed in Master Sheet" if not item.urls else None),
        "Lowest Price Found": lowest, "Source Notes": " | ".join(notes)[:900],
        "Reference Price": ref[0], "Reference Type": f"{ref[1]}: {ref[2]}" if ref[0] else ref[1],
        "Prior Verified Price": hist.get("prior"), "Change vs Prior": change,
        "Verified Low": min([v for v in [hist.get("low")] + [l.unit_price for l in listings if l.trusted and l.pack_qty == 1]
                             if v is not None], default=None),
        "EWMA (Verified)": hist.get("ewma"), "Qty Needed": item.qty_needed,
        "Best Qty Plan": plan["text"][:400] if plan else "no purchasable High-confidence offer",
        "Best Qty Total": plan["total"] if plan else None,
        "Product Identity": " | ".join(identity) or "title/spec matching only (no identifiers known yet)",
        "Evidence Mix": evidence_mix or "no eligible listings",
        "Retrieval Outcomes": _outcome_summary(results)[:600],
        "Search Mode": mode, "Primary Link Failed": link_flag[:300], "Expanded Search No Results": none_flag[:300],
    }
    deal_rows = []
    for l in deals:
        r = l.ref_price
        deal_rows.append({
            "runID": ctx.run_id, "WishlistItem": item.wid, "Product": item.product, "URL": l.url,
            "isPrimaryVendor": 1 if l.is_primary else 0, "Target Price": tgt, "Price": l.unit_price,
            "RightProductConfidence": l.confidence, "Vendor": l.vendor, "Source": l.source,
            "Condition": l.condition, "Listed Price": l.price, "Pack Qty": l.pack_qty,
            "% Below Target": round((tgt - l.unit_price) / tgt, 4) if tgt else None,
            "Baseline Price": r, "% Below Baseline": round((r - l.unit_price) / r, 4) if r else None,
            "Deal Rule": l.deal_rule, "Secondary Vendor?": "Yes" if l.is_resale else "No",
            "Secondary Vendor Comments": (l.seller_comment or "N/A (no seller rating available)") if l.is_resale else "",
            "In Stock Verified": stock_status(l), "Listing Title": l.title[:200],
            "Evidence": l.evidence, "Match Evidence": f"{l.match_evidence or 'title'}: {l.conf_reason}"[:200],
            "Reference Type": l.ref_type, "Regular Price": l.regular_price, "Shipping": l.shipping,
            "Availability": l.availability, "Conditional Pricing": (f"{l.conditional}: {l.conditional_detail}"
                                                                    if l.conditional else ""),
            "Corroborated By": "; ".join(l.corroborated_by)[:300], "Price Event": l.price_event,
            "Offer ID": l.offer_id, "Retrieved At": (l.retrieved_at or ctx.now).strftime("%Y-%m-%d %H:%M")})
    return run_row, deal_rows, deals


def stock_status(l: Listing) -> str:
    """'In Stock Verified': only merchant pages / official APIs can confirm stock."""
    if l.evidence in VERIFIED:
        return "Yes" if l.in_stock else ("No" if l.in_stock is False else "N/A")
    return "N/A"


# =============================================================================
# 4. WORKBOOK WRITER  (touches ONLY Run Data and Deals Data)
# =============================================================================

def ensure_columns(ws, wanted: list) -> dict:
    hm = header_map(ws)
    last = max(hm.values()) if hm else 0
    for name in wanted:
        if norm_header(name) not in hm:
            last += 1
            cell = ws.cell(1, last, name)
            if last > 1:
                src = ws.cell(1, last - 1)
                cell.font, cell.fill = copy(src.font), copy(src.fill)
                cell.alignment, cell.border = copy(src.alignment), copy(src.border)
            hm[norm_header(name)] = last
        if name in COL_WIDTHS:
            letter = get_column_letter(hm[norm_header(name)])
            cur = ws.column_dimensions[letter].width
            if not cur or cur < COL_WIDTHS[name]:
                ws.column_dimensions[letter].width = COL_WIDTHS[name]
    return hm


def next_empty_row(ws) -> int:
    r = ws.max_row
    while r > 1 and all(c.value is None for c in ws[r]):
        r -= 1
    return r + 1


def remove_run_rows(ws, run_id: str, wids: set) -> int:
    """Idempotency: delete rows already written for (runID, WishlistItem) before re-writing them."""
    hm = header_map(ws)
    cr, cw = find_col(hm, "runID"), find_col(hm, "WishlistItem")
    if not cr:
        return 0
    n = 0
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, cr).value or "") == run_id and (not cw or str(ws.cell(r, cw).value).strip() in wids):
            ws.delete_rows(r)
            n += 1
    return n


def append_rows(ws, wanted: list, rows: list) -> None:
    hm = ensure_columns(ws, wanted)
    r = next_empty_row(ws)
    for row in rows:
        for name, val in row.items():
            col = hm.get(norm_header(name))
            if col is None or val is None:
                continue
            cell = ws.cell(r, col, val)
            if name == "RowDateTime":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif name in MONEY_COLS:
                cell.number_format = "$#,##0.00"
            elif name in PCT_COLS:
                cell.number_format = "0.0%"
        r += 1


def write_github_summary(outcomes: list, run_id: str) -> None:
    path = os.getenv("GITHUB_STEP_SUMMARY")
    if not path:
        return
    lines = [f"### Price tracker run `{run_id}`", "",
             "| Item | Verified price | Reference | Deals (new events) | Best plan | Outcomes |", "|---|---|---|---|---|---|"]
    for item, run_row, _, deals in outcomes:
        new_ev = sum(1 for d in deals if not d.price_event.startswith("unchanged"))
        vb = run_row.get("Vendor Baseline (New)")
        rp = run_row.get("Reference Price")
        lines.append(f"| {item.product} | {f'${vb:,.2f}' if vb else '-'} | {f'${rp:,.2f}' if rp else '-'} | "
                     f"{len(deals)} ({new_ev}) | {(run_row.get('Best Qty Plan') or '-')[:80]} | "
                     f"{(run_row.get('Retrieval Outcomes') or '')[:120]} |")
    try:
        with open(path, "a", encoding="utf-8") as fh:
            fh.write("\n".join(lines) + "\n")
    except OSError:
        pass


# =============================================================================
# 5. MAIN
# =============================================================================

def parse_args(argv=None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Home wishlist price tracker")
    p.add_argument("--workbook", help="path to the .xlsx (default: auto-detect Home*Wishlist*.xlsx)")
    p.add_argument("--min-priority", type=int, default=1)
    p.add_argument("--locations", default="")
    p.add_argument("--items", default="")
    p.add_argument("--max-items", type=int, default=0)
    p.add_argument("--force", "--allproducts", dest="force", action="store_true",
                   help="check every selected product now, ignoring CADENCE_DAYS")
    p.add_argument("--no-serpapi", action="store_true")
    p.add_argument("--no-ebay", action="store_true")
    p.add_argument("--no-direct", action="store_true", help="skip merchant pages/APIs (aggregators only)")
    p.add_argument("--no-discovery", action="store_true", help="skip known-retailer discovery (Product URLs/APIs still run)")
    p.add_argument("--force-discovery", action="store_true", help="ignore discovery/SerpApi caches for this run")
    p.add_argument("--browser", choices=("auto", "off"), default=os.getenv("PRICE_TRACKER_BROWSER", "auto"),
                   help="Playwright last-resort renderer (auto = use it if installed)")
    p.add_argument("--workers", type=int, default=int(os.getenv("PRICE_TRACKER_WORKERS", WORKERS)))
    p.add_argument("--run-id", default=os.getenv("PRICE_TRACKER_RUN_ID", ""),
                   help="re-using a run id replaces that run's rows instead of duplicating them")
    p.add_argument("--state-dir", default="", help=f"state folder (default: {STATE_DIR}/ next to the workbook)")
    p.add_argument("--dry-run", action="store_true", help="do everything but do NOT write the workbook or state")
    return p.parse_args(argv)


def find_workbook(arg: Optional[str]) -> Path:
    if arg:
        p = Path(arg)
        if p.exists():
            return p
        raise FileNotFoundError(f"Workbook not found: {arg}")
    for base in (Path("."), Path(__file__).resolve().parent):
        matches = sorted(base.glob(WORKBOOK_GLOB))
        if matches:
            return matches[0]
    raise FileNotFoundError(f"No workbook matching '{WORKBOOK_GLOB}' in {Path('.').resolve()}")


def select_items(items: list, args: argparse.Namespace, history: list, now: datetime) -> list:
    ids = {s.strip() for s in args.items.split(",") if s.strip()}
    locs = [s.lower() for s in split_multi(args.locations)]
    chosen = []
    for it in items:
        if it.priority < args.min_priority:
            continue
        if ids and str(it.wid).strip() not in ids:
            continue
        if locs and not any(l in it.location.lower() for l in locs):
            continue
        if not args.force:
            gap = CADENCE_DAYS.get(it.priority, 0)
            last = max((h["dt"] for h in history if h["wid"] == str(it.wid).strip()), default=None)
            if gap and last and (now - last) < timedelta(days=gap) - timedelta(hours=12):
                log(f"  skip '{it.product}' (priority {it.priority}: last checked {last:%Y-%m-%d}, cadence {gap}d)")
                continue
        chosen.append(it)
    chosen.sort(key=lambda i: -i.priority)
    return chosen[:args.max_items] if args.max_items else chosen


def load_dotenv(path: Path) -> None:
    """Minimal .env reader. Existing environment variables win (GitHub secrets are never overridden)."""
    try:
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            k, v = k.strip().removeprefix("export ").strip(), v.strip().strip('"').strip("'")
            if k and k not in os.environ:
                os.environ[k] = v
    except OSError:
        pass


def build_context(args, wb, run_id, now, history, state, session=None) -> Context:
    primary = load_primary_vendors(wb[SHEET_PRIMARY])
    browser = BrowserRenderer(enabled=(args.browser != "off"))
    fetcher = Fetcher(state.data, session=session or requests.Session(), browser=browser)
    shopify, bestbuy = ShopifyAdapter(), BestBuyAdapter()
    page = ProductPageAdapter(shopify=shopify, bestbuy=bestbuy if bestbuy.available()[0] else None)
    serp_client = SerpApiClient(os.getenv("SERPAPI_KEY"))
    ebay_client = EbayClient(os.getenv("EBAY_CLIENT_ID"), os.getenv("EBAY_CLIENT_SECRET"))
    serp_client.log = ebay_client.log = log
    adapters = {"page": page, "shopify": shopify, "bestbuy": bestbuy,
                "discovery": RetailerDiscoveryAdapter(page, shopify),
                "serpapi": SerpApiAdapter(serp_client), "ebay": EbayAdapter(ebay_client)}
    actx = AdapterContext(fetcher=fetcher, state=state, primary=primary, log=log,
                          no_browser=(args.browser == "off"), force_discovery=args.force_discovery)
    return Context(run_id=run_id, now=now, primary=primary, secondary_keys=load_secondary_keys(wb[SHEET_SECONDARY]),
                   history=history, state=state, fetcher=fetcher, adapters=adapters, actx=actx, args=args)


def main(argv=None) -> int:
    if not os.getenv("PRICE_TRACKER_NO_DOTENV"):          # tests set this so a local .env is never used
        for env_file in (Path(".env"), Path(__file__).resolve().with_name(".env")):
            load_dotenv(env_file)
    args = parse_args(argv)
    _push_settings()
    try:
        path = find_workbook(args.workbook)
        wb = load_workbook(path)
        for name in (SHEET_MASTER, SHEET_PRIMARY, SHEET_SECONDARY, SHEET_RUN, SHEET_DEALS):
            if name not in wb.sheetnames:
                raise KeyError(f"Missing worksheet '{name}'")
    except Exception as e:
        log(f"FATAL: cannot open workbook: {e}")
        return 2

    now = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
    run_id = args.run_id.strip() or (now.strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:4])
    state = State(Path(args.state_dir) if args.state_dir else path.resolve().parent / STATE_DIR, now)
    log(f"Run {run_id} | workbook: {path} | dry-run: {args.dry_run}")

    items = load_items(wb[SHEET_MASTER])
    history = load_history(wb[SHEET_RUN])
    todo = select_items(items, args, [h for h in history if h.get("run") != run_id] if args.run_id else history, now)
    log(f"{len(items)} wishlist items, {len(todo)} selected for this run")
    if not todo:
        return 0

    ctx = build_context(args, wb, run_id, now, history, state)
    serp = ctx.adapters["serpapi"].client
    if not args.no_serpapi:
        if serp.disabled:
            log("NOTE: SERPAPI_KEY is not set - Google Shopping discovery is OFF (optional; merchant pages/APIs still run).")
        serp.check_credits()
    if not ctx.adapters["bestbuy"].available()[0]:
        log(f"NOTE: Best Buy API off - {ctx.adapters['bestbuy'].available()[1]}")
    if not args.no_ebay and ctx.adapters["ebay"].client.disabled and any(i.open_used for i in todo):
        log("NOTE: eBay keys not set - eBay is optional and skipped.")
    if args.browser != "off" and not ctx.fetcher.browser.enabled:
        log(f"NOTE: browser fallback off - {ctx.fetcher.browser.reason}")

    outcomes = []
    try:
        for it in todo:
            log(f"\n> [{it.wid}] {it.product} (priority {it.priority}, target ${it.target}, qty {it.qty_needed}, "
                f"{'Product URLs only (+ same-vendor fallback)' if it.only_links else 'wide+used' if it.open_used else 'primary vendors only'}"
                f"{', excluding ' + '; '.join(it.exclude) if it.exclude else ''})")
            try:
                run_row, deal_rows, deals = process_item(it, ctx)
            except Exception as e:     # one bad item must never sink the whole run
                log(f"  ERROR processing item: {type(e).__name__}: {e}")
                run_row = {"RowDateTime": now, "runID": run_id, "WishlistItem": it.wid, "Product": it.product,
                           "Listings Searched": 0, "Unique Websites Searched": 0,
                           "Target Price or better found": 0, "Primary Vendor Price or Better": 0,
                           "Deals found": 0, "Primary Vendor Deals": 0, "Qty Needed": it.qty_needed,
                           "Source Notes": f"ERROR: {type(e).__name__}: {e}"[:300],
                           "Retrieval Outcomes": f"{Outcome.PARSER} 1 (item processing)"}
                deal_rows, deals = [], []
            log(f"  listings={run_row['Listings Searched']} verified=${run_row.get('Vendor Baseline (New)')} "
                f"ref=${run_row.get('Reference Price')} deals={run_row['Deals found']}")
            log(f"  outcomes: {run_row.get('Retrieval Outcomes')}")
            log(f"  notes: {run_row['Source Notes']}")
            if run_row.get("Best Qty Plan"):
                log(f"  qty plan: {run_row['Best Qty Plan']}")
            for d in deals:
                log(f"  DEAL ${d.unit_price:,.2f} @ {d.vendor} [{d.confidence}/{d.evidence}] {d.deal_rule} "
                    f"({d.price_event}) -> {d.url}")
            outcomes.append((it, run_row, deal_rows, deals))
    finally:
        ctx.fetcher.browser and ctx.fetcher.browser.close()

    if args.dry_run:
        log("\nDry run: workbook and state NOT modified.")
    else:
        try:
            wids = {str(o[0].wid).strip() for o in outcomes}
            removed = remove_run_rows(wb[SHEET_RUN], run_id, wids) + remove_run_rows(wb[SHEET_DEALS], run_id, wids)
            if removed:
                log(f"Re-run of {run_id}: replaced {removed} previously written row(s).")
            append_rows(wb[SHEET_RUN], RUN_COLS, [o[1] for o in outcomes])
            append_rows(wb[SHEET_DEALS], DEALS_COLS, [r for o in outcomes for r in o[2]])
            tmp = path.with_name(path.name + ".tmp")
            wb.save(tmp)
            os.replace(tmp, path)
            state.save()
            log(f"\nSaved {len(outcomes)} Run Data rows and {sum(len(o[2]) for o in outcomes)} Deals Data rows; "
                f"state in {state.folder}.")
        except Exception as e:
            log(f"FATAL: could not save workbook: {type(e).__name__}: {e}")
            return 2
    write_github_summary(outcomes, run_id)
    log(f"HTTP: {ctx.fetcher.stats}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
