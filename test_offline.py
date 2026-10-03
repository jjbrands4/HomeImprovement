#!/usr/bin/env python3
"""
test_offline.py - verifies price_tracker.py + pricetrack/ end-to-end WITHOUT internet or API keys.

A fake web (merchant pages, Shopify JSON, Best Buy API, SerpApi, eBay) is routed through the real
fetcher, adapters, matching, reconciliation, pricing, history and workbook writer. Checks cover:
  [1] identity: identifier-first matching, variant conflicts (region/generation/colour/voltage/size/
      accessory/bundle/sibling model), Product URL validation
  [2] URL normalisation, GTINs, structured-data extraction (all offers/variants, conditional prices)
  [3] resilient fetching: backoff/retry, 403 handling, circuit breaker, conditional GET cache
  [4] adapters: Shopify variants, Best Buy identifier search, discovery cache, stale redirects,
      browser last-resort fallback
  [5] pricing: quantity plan, reference-price hierarchy, conditional pricing kept out of averages
  [6] end-to-end runs: evidence tiers (aggregator fallback never verified), deals, provenance columns,
      offer change events, idempotent re-runs, dry-run, untouched worksheets, no keys / no network
  [7] Master Sheet switches: ';' lists, '!' exclusions, 'Only Check Primary Links' (Product URLs only +
      same-vendor SerpApi fallback)

Run:  python test_offline.py            (from the repo folder containing the workbook)
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import urlencode

from openpyxl import load_workbook

import price_tracker as pt
import pricetrack.fetch as pfetch
from pricetrack import identity as idn, pricing, reconcile
from pricetrack.adapters import (AdapterContext, BestBuyAdapter, ProductPageAdapter, RetailerDiscoveryAdapter,
                                 ShopifyAdapter, parse_serpapi)
from pricetrack.extract import extract_product
from pricetrack.fetch import Fetcher
from pricetrack.history import State, history_stats, verified_series
from pricetrack.models import Item, Listing, Outcome, Vendor, MARKET_SNAPSHOT, VERIFIED_DIRECT, VERIFIED_DISCOVERED
from pricetrack.urls import merchant_item_id, normalize_url

pfetch.cffi_requests = None                 # never touch the real network from tests
os.environ["PRICE_TRACKER_NO_DOTENV"] = "1"  # never load your real keys from .env during tests
pfetch.time.sleep = lambda *_: None         # no politeness / backoff delays in tests
pt.time.sleep = lambda *_: None

SRC = next(iter(sorted(Path(".").glob(pt.WORKBOOK_GLOB))), None)
if SRC is None:
    sys.exit("Put this file next to your workbook (Home Wishlist.xlsx) and re-run.")

passed = failed = 0


def check(label, cond, extra=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {label}")
    else:
        failed += 1
        print(f"  FAIL  {label}  {extra}")


# ---------------------------------------------------------------------------
# Fake web
# ---------------------------------------------------------------------------
class Resp:
    def __init__(self, status=200, data=None, text="", ctype="text/html", headers=None, url=""):
        self.status_code, self._d = status, data
        self.text = text if data is None else json.dumps(data)
        self.content = self.text.encode()
        self.headers = {"content-type": ctype if data is None else "application/json", **(headers or {})}
        self.ok = status < 400
        self.url = url

    def json(self):
        if self._d is None:
            return json.loads(self.text)
        return self._d


class Web:
    """requests.Session stand-in: routes by URL substring (first match wins), logs every request.
    A route value may be a Resp, an Exception, a list (consumed in order) or a callable(url, headers)."""
    def __init__(self, routes=()):
        self.routes, self.log, self.headers_seen = list(routes), [], []

    def get(self, url, params=None, headers=None, **kw):
        full = url + ("?" + urlencode(params) if params else "")
        self.log.append(full)
        self.headers_seen.append(headers or {})
        for pat, resp in self.routes:
            if pat in full:
                if isinstance(resp, list):
                    resp = resp.pop(0) if len(resp) > 1 else resp[0]
                if callable(resp) and not isinstance(resp, Resp):
                    resp = resp(full, headers or {})
                if isinstance(resp, Exception):
                    raise resp
                if not resp.url:
                    resp.url = url
                return resp
        return Resp(404, None, "not found", url=url)
    post = get


def ld(name, price, avail="InStock", gtin=None, mpn=None, regular=None, brand=None):
    offer = {"@type": "Offer", "price": str(price), "priceCurrency": "USD", "availability": f"https://schema.org/{avail}"}
    if regular:
        offer["priceSpecification"] = [{"@type": "UnitPriceSpecification", "price": regular,
                                        "priceType": "https://schema.org/StrikethroughPrice"}]
    node = {"@context": "https://schema.org", "@type": "Product", "name": name, "offers": offer}
    if gtin:
        node["gtin12"] = gtin
    if mpn:
        node["mpn"] = mpn
    if brand:
        node["brand"] = {"@type": "Brand", "name": brand}
    return Resp(200, None, f'<html><title>{name}</title><script type="application/ld+json">{json.dumps(node)}</script>'
                           + "<p>" + "product description words " * 120 + "</p></html>")


def mk_item(product, kws="", specs="", wid=1, qty=1, open_used=True, target=None, bulk=False, sizes=(), urls=()):
    kw = [k.strip() for k in kws.split(";") if k.strip()]
    sp = [s.strip() for s in specs.split(";") if s.strip()]
    skus = [k for k in kw if idn.looks_like_sku(k)]
    return Item(wid, product, sp, kw, "Room", 3, qty, open_used, target, bulk, set(sizes), skus=skus,
                spec_phrases=sp + [k for k in kw if k not in skus], urls=list(urls))


def actx(web, state=None, browser=None):
    st = state or State(None)
    return AdapterContext(fetcher=Fetcher(st.data, session=web, browser=browser), state=st, primary=[])


hue = mk_item("Hue Color Slim Downlight 4-INCH", "609404; 4 inch", target=60)
plug = mk_item("THIRDREALITY Smart Plug Gen3", wid=2, qty=2, open_used=False, target=13, bulk=True, sizes=(2, 4))
sonos = mk_item("Sonos Ray Soundbar", "RAYG1US1BLK", "Unmounted", wid=3, target=175)

# ---------------------------------------------------------------------------
print("\n[1] Identity: identifier-first matching + variant conflicts")
cc = lambda it, t, **kw: idn.classify(it, t, **kw).confidence   # noqa: E731
cases = [
    (hue, 'Philips Hue White and Color Ambiance 4" Slim Downlight 609404', "High", "SKU + 4-inch"),
    (hue, "Philips Hue Color Slim Downlight 4-inch", "High", "exact name"),
    (hue, "Philips Hue White Ambiance Slim Downlight 4 inch", "Low", "White Ambiance is another product"),
    (hue, "Case for Hue Color Slim Downlight 4-inch", "Low", "accessory"),
    (hue, "Philips Hue Color Slim Downlight 6-inch", "Low", "different size"),
    (hue, "Philips Hue Lightstrip Plus", "Low", "sibling / configuration"),
    (hue, "Philips Hue Color Slim Downlight 4 inch, compatible with Hue Bridge", "High", "'compatible with Bridge' is not an accessory"),
    (plug, "ThirdReality Zigbee Smart Plug Gen3, Power Meter, Works with Home Assistant", "High", "inserted word ok"),
    (plug, "THIRDREALITY Smart Plug Gen3 4 Pack, Precise Real-time Power Meter", "High", "pack phrase ignored"),
    (plug, "THIRDREALITY Smart Plug Gen2", "Low", "generation"),
    (plug, "THIRDREALITY Smart Plug E2", "Low", "sibling model E2"),
    (plug, "THIRDREALITY Smart Dual Plug ZP1", "Low", "sibling model"),
    (plug, "Kasa Smart Plug Mini 15A", "Low", "other brand / Mini"),
    (sonos, "Sonos Ray Compact Soundbar Black RAYG1US1BLK", "Medium", "SKU but spec missing"),
    (sonos, "Sonos Ray Soundbar White RAYG1US1WHT", "Low", "colour variant code"),
    (sonos, "Sonos Ray - Compact Smart Soundbar (Black) | RAYG1EU1BLK", "Low", "regional (EU) model code"),
    (sonos, "Sonos Ray Soundbar (International Version)", "Low", "regional wording"),
    (sonos, "Sonos Ray Soundbar 220-240V", "Low", "non-US voltage"),
    (sonos, "Sonos Ray Soundbar 100-240V Unmounted", "High", "universal PSU is not regional"),
    (sonos, "Sonos Ray Soundbar + Sub Mini Bundle", "Low", "bundle"),
    (sonos, "Sonos Ray Soundbar with Sub Mini", "Low", "bundle via 'with'"),
    (sonos, "Sonos Beam (Gen 2) Smart Soundbar Black", "Low", "generic sibling rule (no family table needed)"),
    (sonos, "Mounting Kit Compatible with Sonos Ray Soundbar", "Low", "accessory"),
    (sonos, "Sonos Ray Soundbar - White", "Low", "colour conflicts with SKU colour (BLK)"),
]
for it, title, exp, why in cases:
    got = idn.classify(it, title)
    check(f"{why}: '{title[:52]}' -> {exp}", got.confidence == exp, f"got {got.confidence}: {got.reason}")

g = idn.normalize_gtin("046677609405")
hue_g = mk_item("Hue Color Slim Downlight 4-INCH", "609404; 4 inch")
hue_g.learned_gtins = {g}
check("GTIN match beats a weak title -> High", cc(hue_g, "Philips Downlight", gtins={g}) == "High")
check("Different GTIN on a single unit -> Low", cc(hue_g, "Philips Hue Color Slim Downlight 4 inch",
                                                   gtins={idn.normalize_gtin("046677562434")}) == "Low")
check("Different GTIN on a 4-pack is not a conflict (packs have own GTIN)",
      cc(hue_g, "Philips Hue Color Slim Downlight 4 inch 4 Pack", gtins={idn.normalize_gtin("046677562434")}, pack_qty=4) == "High")
check("Structured MPN match (page metadata) -> High", idn.classify(sonos, "Sonos Ray Unmounted", mpns={"RAYG1US1BLK"}).evidence == "page_metadata")
check("Supplemental family table still active", idn.SUPPLEMENTAL_FAMILY_RULES and cc(sonos, "Sonos Ray Arc combo") == "Low")

vp = idn.validate_page
check("Product URL page 'Sonos Ray' on sonos.com validates (brand implied by domain)",
      vp(sonos, "Sonos Ray", set(), set(), domain_brand="Sonos Direct sonos.com", slug="en us shop ray black").confidence == "High")
check("Product URL page redirected to a different product -> Low",
      vp(sonos, "Sonos Beam (Gen 2)", set(), set(), domain_brand="sonos.com").confidence == "Low")
check("Product URL page with vague title -> Medium (not trusted)",
      vp(hue, "Philips Hue Slim Downlight | Dell USA", set(), set(), domain_brand="Dell dell.com").confidence == "Medium")

# ---------------------------------------------------------------------------
print("\n[2] URLs, GTINs and structured-data extraction")
u = ("https://www.sonos.com/en-us/shop/ray-black?utm_campaign=x&gclid=abc&gad_source=1&gbraid=z#reviews")
check("Tracking params + fragment removed", normalize_url(u) == "https://www.sonos.com/en-us/shop/ray-black", normalize_url(u))
check("Functional variant param kept", normalize_url("https://a.com/products/x?srsltid=1&variant=55") == "https://a.com/products/x?variant=55")
check("Best Buy ref/loc removed, sku path kept",
      normalize_url("https://www.bestbuy.com/product/x/J39/sku/6506474?utm_source=feed&ref=212&loc=1") ==
      "https://www.bestbuy.com/product/x/J39/sku/6506474")
check("Merchant item ids (Best Buy / Amazon / Shopify variant)",
      (merchant_item_id("https://www.bestbuy.com/site/x/6506474.p?skuId=6506474"),
       merchant_item_id("https://www.amazon.com/dp/B0B4F9LTV1?th=1"),
       merchant_item_id("https://s.com/products/p?variant=77")) == ("6506474", "B0B4F9LTV1", "v77"))
check("UPC check digit validated", idn.normalize_gtin("046677609405") == "00046677609405" and idn.normalize_gtin("046677609404") is None)

group = {"@context": "https://schema.org", "@type": "ProductGroup", "name": "THIRDREALITY Smart Plug Gen3",
         "brand": {"name": "THIRDREALITY"}, "hasVariant": [
             {"@type": "Product", "sku": "P1", "gtin13": "0046677609405", "additionalProperty": [{"name": "Pack", "value": "1 Pack"}],
              "offers": {"@type": "Offer", "price": "14.99", "priceCurrency": "USD", "availability": "https://schema.org/OutOfStock"}},
             {"@type": "Product", "sku": "P2", "additionalProperty": [{"name": "Pack", "value": "2"}],
              "offers": {"@type": "Offer", "price": "26.99", "priceCurrency": "USD", "availability": "InStock",
                         "priceSpecification": [{"price": 29.99, "priceType": "https://schema.org/StrikethroughPrice"},
                                                {"price": 24.99, "validForMemberTier": {"name": "Plus"}}]}}]}
pg = extract_product(f'<script type="application/ld+json">{json.dumps(group)}</script>')
check("JSON-LD ProductGroup: every variant extracted (not the first price)", pg and sorted(o.price for o in pg.offers) == [14.99, 26.99],
      str(pg and [(o.price, o.variant) for o in pg.offers]))
o2 = next(o for o in pg.offers if o.price == 26.99)
check("Variant pack size, strikethrough regular price, stock", o2.pack_qty == 2 and o2.regular_price == 29.99 and o2.in_stock is True)
check("Member price recorded separately, NOT used as the price", "member price $24.99" in o2.conditional_detail and o2.price == 26.99)
check("Variant GTIN kept on its own offer", next(o for o in pg.offers if o.price == 14.99).gtin == "00046677609405")
dell_html = '<html><title>Hue | Dell USA</title><div itemscope itemtype="https://schema.org/Product"><span itemprop="price" content="64.99"></span><meta itemprop="gtin12" content="046677609405"></div></html>'
mp = extract_product(dell_html)
check("Microdata price + GTIN", mp and mp.offers[0].price == 64.99 and "00046677609405" in mp.gtins)
meta = '<meta property="og:title" content="Sonos Ray"><meta property="product:price:amount" content="199.00"><meta property="product:availability" content="out of stock">'
check("OpenGraph meta price + availability", extract_product(meta).offers[0].in_stock is False)
hyd = ('<script id="__NEXT_DATA__" type="application/json">{"props":{"accessory":{"price":19.99},'
       '"product":{"name":"Sonos Ray","skuId":"6506474","currentPrice":219.0,"regularPrice":279.0,"memberPrice":199.0}}}</script>')
hp = extract_product(hyd, 175, "Sonos Ray Soundbar")
check("Scoped hydration JSON: product node chosen (not the $19.99 accessory), regular kept, member price ignored",
      hp and hp.offers[0].price == 219.0 and hp.offers[0].regular_price == 279.0 and "membership" in hp.offers[0].conditional_detail,
      str(hp and hp.offers))
check("DOM fallback skips financing '/mo' prices",
      (extract_product('<span class="price">$18.25</span>/mo <div class="product-price">$219.00</div>', 175) or
       type("x", (), {"offers": [type("o", (), {"price": 0})]})).offers[0].price == 219.0)

# ---------------------------------------------------------------------------
print("\n[3] Resilient fetching")
web = Web([("flaky", [Resp(503, None, "busy"), Resp(503, None, "busy"), Resp(200, None, "<html>ok</html>")])])
sleeps = []
f = Fetcher({}, session=web, sleep=sleeps.append)
r = f.get("https://flaky.example.com/p")
check("503 retried with backoff, then success", r.ok and len([u for u in web.log if "flaky" in u]) == 3 and len(sleeps) >= 2, str(web.log))
check("Backoff grows (exponential with jitter)", len(sleeps) >= 2 and sleeps[-1] >= sleeps[0] * 0.5)
web = Web([("forbid", Resp(403, None, "denied"))])
f = Fetcher({}, session=web, sleep=lambda s: None)
r = f.get("https://forbid.example.com/a")
check("403 classified as blocked and NOT retried with the same client", r.outcome == Outcome.BLOCKED and len(web.log) == 1)
f.get("https://forbid.example.com/b"); f.get("https://forbid.example.com/c")
r = f.get("https://forbid.example.com/d")
check("Circuit opens after 3 consecutive blocks - 4th request never sent",
      r.outcome == Outcome.BLOCKED and "circuit" in r.detail and len(web.log) == 3, r.detail)
check("Block cool-down persisted to state", "forbid.example.com" in f.state["circuits"])
web = Web([("captcha", Resp(200, None, "<html>Please verify you are human - px-captcha</html>"))])
check("200 bot-challenge page classified as blocked", Fetcher({}, session=web, sleep=lambda s: None).get("https://captcha.x.com").outcome == Outcome.BLOCKED)
web = Web([("throttle", [Resp(429, None, "", headers={"Retry-After": "7"}), Resp(200, None, "<html>ok</html>")])])
sl = []
Fetcher({}, session=web, sleep=sl.append).get("https://throttle.example.com/")
check("Retry-After honoured", 7.0 in sl, str(sl))


def etag_route(url, headers):
    if headers.get("If-None-Match") == '"v1"':
        return Resp(304, None, "", headers={"ETag": '"v1"'})
    return Resp(200, {"title": "P", "variants": [{"id": 1, "title": "Default Title", "price": 1499, "available": True}]},
                headers={"ETag": '"v1"'})


st = State(None)
web = Web([("/products/p.js", etag_route)])
c1 = actx(web, st)
r1 = ShopifyAdapter().product("https://shop.example.com", "p", plug, c1)
r2 = ShopifyAdapter().product("https://shop.example.com", "p", plug, c1)
check("Conditional GET: second fetch sends If-None-Match and reuses the parsed result on 304",
      r1.listings and r2.listings and r2.listings[0].price == 14.99 and "304" in r2.detail, r2.detail)

# ---------------------------------------------------------------------------
print("\n[4] Adapters")
shop_js = {"title": "THIRDREALITY Smart Plug Gen3", "vendor": "THIRDREALITY", "options": [{"name": "Pack"}],
           "variants": [{"id": 11, "title": "1 Pack", "option1": "1 Pack", "price": 1499, "compare_at_price": 1799,
                         "available": False, "sku": "TR-PLUG-G3", "barcode": "046677609405"},
                        {"id": 22, "title": "2 Pack", "option1": "2 Pack", "price": 2699, "available": True},
                        {"id": 44, "title": "4 Pack", "option1": "4 Pack", "price": 4999, "available": True}]}
web = Web([("/search/suggest.json", Resp(200, {"resources": {"results": {"products": [
              {"title": "THIRDREALITY Smart Plug Gen3", "handle": "smart-plug-gen3"},
              {"title": "THIRDREALITY Smart Plug E2", "handle": "smart-plug-e2"},
              {"title": "THIRDREALITY Motion Sensor", "handle": "motion"}]}}})),
           ("/products/smart-plug-gen3.js", Resp(200, shop_js)),
           ("/products.json", Resp(200, {"products": []}))])
c = actx(web)
sh = ShopifyAdapter()
check("Shopify auto-detected (no manual DTC flag needed)", sh.is_shopify("www.thirdreality.com", c) is True)
r = sh.search(plug, c, domain="www.thirdreality.com", vendor="THIRDREALITY Store")
check("Shopify search keeps only the Gen3 product and reads every variant",
      sorted(l.price for l in r.listings) == [14.99, 26.99, 49.99] and sum(".js" in u and ".json" not in u for u in web.log) == 1, r.detail)
l1 = next(l for l in r.listings if l.price == 14.99)
check("Shopify variant: id, pack, compare-at, barcode GTIN, availability",
      l1.variant_id == "11" and l1.pack_qty_hint == 1 and l1.regular_price == 17.99 and "00046677609405" in l1.gtins
      and l1.in_stock is False and l1.url.endswith("variant=11"))
check("Shopify packs from option values", sorted(l.pack_qty_hint for l in r.listings) == [1, 2, 4])

os.environ.pop("BESTBUY_API_KEY", None)
bb_rows = {"upc=": Resp(200, {"products": []}),
           "modelNumber=RAYG1US1BLK": Resp(200, {"products": [{"sku": 6506474, "name": "Sonos - Ray Soundbar - Black",
                                                              "upc": "878269009993", "modelNumber": "RAYG1US1BLK",
                                                              "salePrice": 179.0, "regularPrice": 219.0,
                                                              "onlineAvailability": True, "freeShipping": True,
                                                              "url": "https://www.bestbuy.com/site/x/6506474.p?skuId=6506474&cmp=RMX"}]})}
web = Web(list(bb_rows.items()))
bb = BestBuyAdapter(key="k")
st = State(None)
r = bb.search(sonos, actx(web, st))
check("Best Buy API resolves by model number (not just SKU URLs)", r.ok if hasattr(r, "ok") else r.outcome == Outcome.SUCCESS and r.listings
      and r.listings[0].price == 179.0, r.detail)
lb = r.listings[0]
check("Best Buy listing: regular price, free shipping, UPC, verified_direct",
      lb.regular_price == 219.0 and lb.shipping == 0.0 and lb.gtins and lb.evidence == VERIFIED_DIRECT)
check("Best Buy SKU cached for the next run", (st.discovery_get(sonos, "bestbuy.com") or {}).get("item_id") == "6506474")
check("No BESTBUY_API_KEY -> api_unavailable outcome, no request", BestBuyAdapter(key="").search(sonos, actx(Web())).outcome == Outcome.API)

# discovery: sitemap -> page, then cached URL reused without re-crawling
sonos_v = Vendor("Sonos Direct", "sonos", True, "www.sonos.com")
web = Web([("/products.json", Resp(404)), ("www.sonos.com/\x00", Resp(404)),
           ("robots.txt", Resp(200, None, "Sitemap: https://www.sonos.com/sitemap-index.xml")),
           ("sitemap-index.xml", Resp(200, None, "<sitemapindex><sitemap><loc>https://www.sonos.com/en-gb/sitemap-products.xml</loc></sitemap>"
                                               "<sitemap><loc>https://www.sonos.com/en-us/sitemap-products.xml</loc></sitemap></sitemapindex>")),
           ("en-us/sitemap-products.xml", Resp(200, None, "<urlset><url><loc>https://www.sonos.com/en-us/shop/ray-wall-mount</loc></url>"
                                                         "<url><loc>https://www.sonos.com/en-us/shop/arc</loc></url>"
                                                         "<url><loc>https://www.sonos.com/en-us/shop/ray</loc></url></urlset>")),
           ("/en-us/shop/ray", ld("Sonos Ray Soundbar - Black", "219.00", mpn="RAYG1US1BLK"))])
st = State(None)
c = actx(web, st)
page = ProductPageAdapter(shopify=ShopifyAdapter())
disc = RetailerDiscoveryAdapter(page, ShopifyAdapter())
r = disc.search(sonos, c, vendor=sonos_v, allow_engines=True)
check("Discovery via sitemap -> verified_discovered listing", r.listings and r.listings[0].evidence == VERIFIED_DISCOVERED
      and r.listings[0].price == 219.0, r.note())
check("en-gb sitemap, wall mount and Arc pages never fetched",
      not any(x in u for u in web.log for x in ("en-gb", "ray-wall-mount", "/shop/arc")), str(web.log))
n_before = len(web.log)
disc2 = RetailerDiscoveryAdapter(page, ShopifyAdapter())          # fresh adapter = next run
r = disc2.search(sonos, c, vendor=sonos_v, allow_engines=True)
check("Next run: cached page re-verified, no sitemap/robots requests",
      r.listings and "cached page re-verified" in r.detail and not any("sitemap" in u or "robots" in u for u in web.log[n_before:]),
      str(web.log[n_before:]))
hd = Vendor("Home Depot", "homedepot", False, "www.homedepot.com")
web2 = Web([("/products.json", Resp(404)), ("homedepot.com/s/", Resp(200, None, "<html>no results</html>")),
            ("robots.txt", Resp(404))])
c2 = actx(web2, st)
r = disc.search(sonos, c2, vendor=hd)
r2 = disc.search(sonos, c2, vendor=hd)
check("Retailer without the product: negative result cached (not re-crawled every run)",
      r.outcome == Outcome.NO_MATCH and r2.outcome == Outcome.SKIPPED and "re-check" in r2.detail, r2.detail)

# stale redirect + browser fallback
web = Web([("old-ray", Resp(200, None, "<html>" + "category words " * 300 + "</html>", url="https://www.sonos.com/en-us/speakers"))])
r = ProductPageAdapter().fetch("https://www.sonos.com/en-us/shop/old-ray", sonos, actx(web), vendor="Sonos Direct")
check("Stale Product URL redirected to a category page -> identity_mismatch", r.outcome == Outcome.IDENTITY_MISMATCH, r.note())


class FakeBrowser:
    enabled, reason, renders = True, "", 0

    def render(self, url):
        FakeBrowser.renders += 1
        return ld("Sonos Ray Soundbar", "219.00").text, url, 200

    def close(self):
        pass


web = Web([("js-page", Resp(200, None, '<html><div id="root"></div><script src="app.js"></script></html>'))])
r = ProductPageAdapter().fetch("https://shop.example.com/js-page/p/123", sonos, actx(web, browser=FakeBrowser()), vendor="X")
check("JS-only page -> Playwright last resort renders and prices it", r.listings and r.listings[0].price == 219.0 and FakeBrowser.renders == 1,
      r.note())
web = Web([("static", ld("Sonos Ray Soundbar", "219.00"))])
FakeBrowser.renders = 0
ProductPageAdapter().fetch("https://shop.example.com/static/p/1", sonos, actx(web, browser=FakeBrowser()), vendor="X")
check("Static page with JSON-LD -> browser NOT used", FakeBrowser.renders == 0)
web = Web([("js-page", Resp(200, None, '<html><div id="root"></div></html>'))])
r = ProductPageAdapter().fetch("https://shop.example.com/js-page/p/9", sonos, actx(web), vendor="X")
check("JS-only page without a browser -> parser_failure with a clear reason",
      r.outcome == Outcome.PARSER and "JavaScript" in r.detail, r.detail)

# ---------------------------------------------------------------------------
print("\n[5] Pricing: quantity plan, reference price, conditional prices")


def L(title, price, vendor="Shop", ev=VERIFIED_DIRECT, pack=None, stock=True, regular=None, ship=None, cond="", src="page"):
    return Listing(title, f"https://{vendor.lower().replace(' ', '')}.com/p/{abs(hash((title, price))) % 10**6}", price, vendor,
                   src, evidence=ev, in_stock=stock, pack_qty_hint=pack, regular_price=regular, shipping=ship, conditional=cond)


prim = [Vendor("THIRDREALITY Store", "thirdreality", True, "thirdreality.com"), Vendor("Amazon", "amazon"), Vendor("Best Buy", "bestbuy")]
for qty, exp in ((4, 49.99), (3, 41.98), (2, 26.99), (1, 14.99)):
    it = mk_item("THIRDREALITY Smart Plug Gen3", wid=2, qty=qty, open_used=False, target=13, bulk=True, sizes=(2, 4))
    ls = [L("THIRDREALITY Smart Plug Gen3", 14.99, "THIRDREALITY Store", pack=1),
          L("THIRDREALITY Smart Plug Gen3 - 2 Pack", 26.99, "THIRDREALITY Store", pack=2),
          L("THIRDREALITY Smart Plug Gen3 - 4 Pack", 49.99, "THIRDREALITY Store", pack=4),
          L("THIRDREALITY Smart Plug Gen3 6 Pack", 54.00, "Amazon", ev=MARKET_SNAPSHOT)]           # 6 not allowed
    pricing.score_listings(it, ls, prim, [])
    plan = pricing.quantity_plan(it, ls)
    check(f"Qty {qty}: least-cost valid combination = ${exp}", plan and plan["total"] == exp, plan and plan["text"])
it = mk_item("THIRDREALITY Smart Plug Gen3", wid=2, qty=4, open_used=False, target=13, bulk=True, sizes=(2, 4))
ls = [L("THIRDREALITY Smart Plug Gen3", 9.00, "Amazon", ev=MARKET_SNAPSHOT, ship=6.0),
      L("THIRDREALITY Smart Plug Gen3 - 4 Pack", 49.99, "THIRDREALITY Store", pack=4)]
pricing.score_listings(it, ls, prim, [])
check("Effective price includes known shipping (9.00 + 6.00 = 15.00/unit)", ls[0].effective_price == 15.0 and ls[0].unit_price == 15.0)
check("Plan uses effective price (4-pack $49.99 beats 4 x $15.00)", pricing.quantity_plan(it, ls)["total"] == 49.99)

ls = [L("Sonos Ray Soundbar Unmounted RAYG1US1BLK", 179.0, "Best Buy", regular=219.0),
      L("Sonos Ray Soundbar Unmounted RAYG1US1BLK", 199.0, "Sonos Direct"),
      L("Sonos Ray Soundbar Unmounted RAYG1US1BLK", 149.0, "Amazon", ev=MARKET_SNAPSHOT, cond="coupon")]
pricing.score_listings(sonos, ls, prim + [Vendor("Sonos Direct", "sonos", True)], [])
ref = pricing.reference_price(sonos, ls, {}, {"sonos"})
check("Reference = published regular/compare-at price (MSRP tier), not the sale price", ref[0] == 219.0 and ref[1] == "msrp_regular", str(ref))
check("Coupon price is NOT eligible for ordinary averages but kept as conditional", not ls[2].eligible and ls[2].cond_ok)
ref2 = pricing.reference_price(sonos, ls[1:2], {}, {"sonos"})
check("No regular published -> manufacturer (DTC) price is the reference", ref2[0] == 199.0 and "manufacturer" in ref2[2], str(ref2))
ref3 = pricing.reference_price(sonos, [], {"median": 205.0, "median_n": 4}, set())
check("No verified offers -> trusted historical reference", ref3[0] == 205.0 and ref3[1] == "historical_verified")

# ---------------------------------------------------------------------------
print("\n[6] Reconciliation")
ls = [L("Sonos Ray Soundbar Unmounted RAYG1US1BLK", 219.0, "Best Buy", src="bestbuy_api"),
      L("Sonos Ray Soundbar Unmounted RAYG1US1BLK", 219.0, "Best Buy", ev=MARKET_SNAPSHOT, src="serpapi"),
      L("Sonos Ray Soundbar Unmounted", 219.0, "Target", ev=MARKET_SNAPSHOT, src="serpapi"),
      L("Sonos Ray Soundbar Unmounted", 219.0, "Walmart", ev=MARKET_SNAPSHOT, src="serpapi")]
pricing.score_listings(sonos, ls, prim, [])
reconcile.assign_ids(sonos, ls)
merged = reconcile.merge_duplicates(ls)
check("Same-merchant Google row merges into the verified API listing as corroboration",
      len(merged) == 3 and any("serpapi" in c for c in merged[0].corroborated_by), str([m.vendor for m in merged]))
check("Different merchants at the SAME price are NOT de-duplicated", {m.vendor for m in merged} == {"Best Buy", "Target", "Walmart"})
med = [L("Sonos Ray Compact Soundbar RAYG1US1BLK", 205.0, "Crutchfield", ev=MARKET_SNAPSHOT, src="serpapi"),
       L("Sonos Ray Soundbar Unmounted RAYG1US1BLK", 209.0, "Best Buy")]
pricing.score_listings(sonos, med, prim, [])
check("Medium before corroboration", med[0].confidence == "Medium")
reconcile.promote(sonos, med)
check("Medium + identifier + independent verified price agreement -> promoted High", med[0].confidence == "High" and "promoted" in med[0].conf_reason)
lone = [L("Sonos Ray Compact Soundbar RAYG1US1BLK", 205.0, "Crutchfield", ev=MARKET_SNAPSHOT, src="serpapi")]
pricing.score_listings(sonos, lone, prim, [])
reconcile.promote(sonos, lone)
check("Medium with no independent evidence stays Medium", lone[0].confidence == "Medium")

# ---------------------------------------------------------------------------
print("\n[7] History: verified-only statistics, idempotent offer events")
now = datetime(2026, 10, 1)
rows = [{"wid": "3", "dt": now - timedelta(days=d), "run": f"r{d}", "baseline": v, "baseline_src": src}
        for d, v, src in ((25, 219.0, "verified: Sonos Direct (page) (median of 1)"),
                          (15, 199.0, "verified: Sonos Direct (page), Best Buy (API) (median of 2)"),
                          (6, 209.0, "verified: Best Buy (API) (median of 1)"),
                          (3, 99.0, "none (no verified merchant page/API priced the product)"),
                          (2, 150.0, "Product URLs: Sonos Direct (median of 1)"))]
ser = verified_series(rows, 3)
check("Only 'verified:' baselines enter history (unverified + legacy rows excluded)", [v for _, v in ser] == [219.0, 199.0, 209.0])
hs = history_stats(ser, [], now)
check("prior / 30d verified median / low / EWMA", hs["prior"] == 209.0 and hs["median"] == 209.0 and hs["low"] == 199.0
      and hs["ewma"] and 199 < hs["ewma"] < 219, str(hs))
st = State(None)
lo = L("Sonos Ray Soundbar", 199.0, "Best Buy")
pricing.normalize_prices(lo)
lo.offer_id = "abc"
e1 = st.offer_event(lo, sonos, "run1")
e1b = st.offer_event(lo, sonos, "run1")                   # retry of the SAME run
st.now = now + timedelta(days=3)
e2 = st.offer_event(lo, sonos, "run2")
lo.price = 189.0
pricing.normalize_prices(lo)
e3 = st.offer_event(lo, sonos, "run3")
check("Offer events: new -> (re-run identical) -> unchanged -> drop",
      e1 == "new offer" and e1b == "new offer" and e2.startswith("unchanged since") and e3.startswith("price drop from $199"),
      str((e1, e1b, e2, e3)))

# ---------------------------------------------------------------------------
print("\n[8] End-to-end: evidence tiers, deals, provenance, idempotency")
fx_dir = Path(tempfile.mkdtemp())
FIXTURE = fx_dir / "Home Wishlist.xlsx"
shutil.copy(SRC, FIXTURE)
wbf = load_workbook(FIXTURE)
ms = wbf[pt.SHEET_MASTER]
mh = pt.ensure_columns(ms, ["Product URLs"])
for r in range(ms.max_row, 1, -1):
    ms.delete_rows(r)
fixture_rows = [
    {"wishlistitem": 1, "product": "Hue Color Slim Downlight 4-INCH", "searchkeywords": "609404; 4 inch",
     "location": "Master Bedroom", "priority": 5, "quantityneeded": 4, "opentoused": "Yes", "targetprice": 60, "isbulkoption": "No",
     "producturls": "https://www.philips-hue.com/en-us/p/hue-white-and-color-ambiance-slim-downlight-4-inch/046677609405; "
                    "https://www.dell.com/en-us/shop/hue-color-slim-downlight-4-inch/apd/ad897711/home-automation"},
    {"wishlistitem": 2, "product": "THIRDREALITY Smart Plug Gen3", "location": "Overall", "priority": 4, "quantityneeded": 2,
     "opentoused": "No", "targetprice": 13, "isbulkoption": "Yes", "bulkkeywords": "2Pack; 4 Pack",
     "producturls": "https://www.thirdreality.com/products/smart-plug-gen3?srsltid=AU7gw4V9"},
    {"wishlistitem": 3, "product": "Sonos Ray Soundbar", "productspecifications": "Unmounted", "searchkeywords": "RAYG1US1BLK",
     "location": "Master Bedroom", "priority": 2, "quantityneeded": 1, "opentoused": "Yes", "targetprice": 175, "isbulkoption": "No",
     "producturls": "https://www.sonos.com/en-us/shop/ray-black?utm_campaign=x&gclid=y; "
                    "https://www.bestbuy.com/product/sonos-ray/J39H373Y64/sku/6506474?utm_source=feed&ref=212"},
]
mh = pt.header_map(ms)
for i, row in enumerate(fixture_rows, start=2):
    for k, v in row.items():
        col = pt.find_col(mh, k)
        if col:
            ms.cell(i, col, v)
wbf.save(FIXTURE)


def S(title, price, source, **kw):
    return {"title": title, "extracted_price": price, "source": source,
            "product_link": "https://www.google.com/shopping/product/123", **kw}


SERP = {
    "sonos": [S("Sonos Ray Soundbar RAYG1US1BLK", 287.42, "Zaytoun"),
              S("Sonos Ray - Compact Smart Soundbar (Black) | RAYG1EU1BLK", 199.0, "Zaytoun Intl"),
              S("Sonos Ray Soundbar + Sub Mini Bundle", 699.00, "Amazon.com"),
              S("Sonos Ray Soundbar Black Unmounted RAYG1US1BLK", 219.00, "Best Buy"),
              S("Sonos Ray Soundbar Black Unmounted RAYG1US1BLK", 219.00, "Sonos"),
              S("Sonos Ray Compact Soundbar Unmounted RAYG1US1BLK", 179.00, "Target"),
              S("Sonos Ray Soundbar Unmounted RAYG1US1BLK", 169.00, "Walmart", extensions=["Clip coupon"])],
    "hue": [S('Philips Hue White and Color Ambiance 4" Slim Downlight 609404', 69.99, "Amazon.com"),
            S("Philips Hue Color Slim Downlight 4-inch", 64.99, "Best Buy"),
            S("Philips Hue Color Slim Downlight 4-inch", 49.99, "B&H", delivery="Free delivery"),
            S("Philips Hue Color Slim Downlight 4-inch (4 Pack)", 250.00, "Amazon.com"),
            S("Case for Hue Color Slim Downlight 4-inch", 12.00, "Etsy"),
            S("Philips Hue White Ambiance Slim Downlight 4 inch", 40.00, "Best Buy"),
            S("Philips Hue Color Slim Downlight 6-inch", 59.00, "Lowe's")],
    "thirdreality": [S("THIRDREALITY Smart Plug Gen3 4 Pack, Precise Real-time Power Meter", 39.96, "Amazon.com"),
                     S("THIRDREALITY Smart Plug Gen2", 9.00, "Amazon.com")],
}


def E(title, price, cond, pct, cnt, user="seller1"):
    return {"title": title, "price": {"value": str(price), "currency": "USD"}, "condition": cond, "itemId": f"v1|{abs(hash(title + str(price)))}|0",
            "itemWebUrl": f"https://www.ebay.com/itm/{abs(hash(title + str(price))) % 10**10}", "seller": {"username": user,
            "feedbackPercentage": str(pct), "feedbackScore": cnt},
            "shippingOptions": [{"shippingCostType": "FIXED", "shippingCost": {"value": "5.00", "currency": "USD"}}]}


EBAY = {"sonos": [E("Sonos Beam (Gen 2) Soundbar Black", 199.99, "Used", 99.5, 900)],
        "hue": [E("Philips Hue Color Slim Downlight 4-inch", 29.99, "Open box", 98.5, 1200, "gooddeals"),
                E("Philips Hue Color Slim Downlight 4-inch", 25.00, "Used", 80.0, 40, "shadyseller")]}


def which(q):
    q = q.lower()
    return next(k for k in ("thirdreality", "sonos", "hue") if k in q)


def serp_route(url, headers):
    from urllib.parse import parse_qs, urlparse as up
    q = parse_qs(up(url).query).get("q", [""])[0]
    return Resp(200, {"shopping_results": SERP[which(q)]})


def ebay_route(url, headers):
    from urllib.parse import parse_qs, urlparse as up
    q = parse_qs(up(url).query).get("q", ["hue"])[0]
    return Resp(200, {"itemSummaries": EBAY.get(which(q), [])})


hue_page = ld("Hue White and color ambiance Slim downlight 4 inch", "69.99", "OutOfStock", gtin="046677609405", brand="Philips Hue")
dell_page = Resp(200, None, '<html><title>Philips Hue White and Color Ambiance Slim Downlight 4 inch | Dell USA</title>'
                            '<span itemprop="price" content="64.99"></span>' + "<p>words </p>" * 300 + '</html>')
WEB_ROUTES = [
    ("serpapi.com/account", Resp(200, {"total_searches_left": 90})),
    ("serpapi.com/search", serp_route),
    ("api.ebay.com/identity", Resp(200, {"access_token": "t"})),
    ("api.ebay.com/buy", ebay_route),
    ("philips-hue.com/en-us/p/", hue_page),
    ("dell.com/en-us/shop/hue", dell_page),
    ("thirdreality.com/products/smart-plug-gen3.js", Resp(200, shop_js)),
    ("sonos.com/en-us/shop/ray-black", ld("Sonos Ray", "219.00", brand="Sonos")),
    ("bestbuy.com", ConnectionError("reset by peer")),
]
fake_web = Web(WEB_ROUTES)
_orig_session = pt.requests.Session
pt.requests.Session = lambda: fake_web
import pricetrack.adapters.serpapi as _ps, pricetrack.adapters.ebay as _pe   # noqa: E401
_ps.requests.Session = _pe.requests.Session = lambda: fake_web
os.environ.update({"SERPAPI_KEY": "dummy", "EBAY_CLIENT_ID": "dummy", "EBAY_CLIENT_SECRET": "dummy"})
os.environ.pop("BESTBUY_API_KEY", None)

wb8 = fx_dir / "run" / "Home Wishlist.xlsx"
wb8.parent.mkdir()
shutil.copy(FIXTURE, wb8)
before = load_workbook(wb8)
rc = pt.main(["--workbook", str(wb8), "--force", "--browser", "off", "--workers", "2", "--run-id", "test-run-1"])
check("main() exit code 0", rc == 0)
w = load_workbook(wb8)
rw, dw = w[pt.SHEET_RUN], w[pt.SHEET_DEALS]
rh, dh = pt.header_map(rw), pt.header_map(dw)
runs = [{k: rw.cell(r, c).value for k, c in rh.items()} for r in range(2, rw.max_row + 1)]
new = {str(r["wishlistitem"]): r for r in runs if r.get("runid") == "test-run-1"}
deals = [{k: dw.cell(r, c).value for k, c in dh.items()} for r in range(2, dw.max_row + 1)]
deals = [d for d in deals if d.get("runid") == "test-run-1"]
for d in deals:
    print(f"     deal: item {d['wishlistitem']} ${d['price']} {d['vendor']} [{d['evidence']}] {d['dealrule']} | {d['priceevent']}")
s3, h1, p2 = new.get("3", {}), new.get("1", {}), new.get("2", {})
for k in ("1", "2", "3"):
    r = new.get(k, {})
    print(f"     item {k}: base={r.get('vendorbaselinenew')} ref={r.get('referenceprice')} [{r.get('referencetype')}] "
          f"avg={r.get('avgpricenew')} | {r.get('retrievaloutcomes')}")
check("3 Run Data rows written", len(new) == 3, str(list(new)))
check("Sonos: verified baseline $219 from sonos.com ONLY (Best Buy Google fallback not verified)",
      s3.get("vendorbaselinenew") == 219.0 and "Best Buy" not in (s3.get("baselinesource") or ""), s3.get("baselinesource"))
check("Sonos: failed Best Buy URL covered by Google row, labelled market_snapshot",
      "via Google Shopping fallback (market_snapshot, not verified)" in (s3.get("producturlstatus") or ""), s3.get("producturlstatus"))
check("Sonos: Product URL status shows timeout/network category for Best Buy",
      "[timeout_network]" in (s3.get("producturlstatus") or ""))
check("Sonos: tracking params stripped from stored URLs", "gclid" not in (s3.get("producturlstatus") or "") and
      all("utm_" not in (d.get("url") or "") for d in deals))
sd = [d for d in deals if str(d["wishlistitem"]) == "3"]
check("Sonos: Target $179 is a deal vs the $219 reference, flagged as market snapshot",
      any(d["price"] == 179.0 and d["baselineprice"] == 219.0 and d["evidence"] == MARKET_SNAPSHOT
          and "market snapshot" in d["dealrule"] for d in sd), str(sd))
check("Sonos: regional RAYG1EU1BLK $199, $287 reseller, bundle and eBay Beam are NOT deals",
      all(d["price"] not in (199.0, 287.42, 699.0, 204.99, 199.99) for d in sd))
check("Sonos: coupon price reported only as CONDITIONAL", all(("CONDITIONAL" in d["dealrule"]) == (d["price"] == 169.0) for d in sd)
      and any(d["price"] == 169.0 and "coupon" in (d["conditionalpricing"] or "") for d in sd), str([(d['price'], d['dealrule']) for d in sd]))
check("Sonos: Sonos' own Google row merged as corroboration (not double counted)",
      "merged as corroboration" in (s3.get("sourcenotes") or ""), s3.get("sourcenotes"))
check("Sonos: retrieval outcomes are categorised", "success" in (s3.get("retrievaloutcomes") or "")
      and "timeout_network" in (s3.get("retrievaloutcomes") or ""), s3.get("retrievaloutcomes"))
check("Sonos: deal rows carry provenance (evidence, match evidence, offer id, event, retrieved at)",
      sd and all(d["evidence"] and d["matchevidence"] and d["offerid"] and d["priceevent"] and d["retrievedat"] for d in sd))

check("Hue: GTIN learned from the validated Product URL page", "GTIN 46677609405" in (h1.get("productidentity") or ""),
      h1.get("productidentity"))
check("Hue: verified baseline from Dell in-stock page ($64.99); sold-out Hue page still informs reference",
      h1.get("vendorbaselinenew") == 64.99, str(h1.get("vendorbaselinenew")) + " " + str(h1.get("baselinesource")))
hd_ = [d for d in deals if str(d["wishlistitem"]) == "1"]
check("Hue: B&H $49.99 (free delivery) is a deal; 6-inch / White Ambiance / case / 4-pack are not",
      any(d["price"] == 49.99 and d["shipping"] == 0 for d in hd_) and all(d["price"] not in (59.0, 40.0, 12.0, 62.5) for d in hd_),
      str([(d['price'], d['vendor']) for d in hd_]))
check("Hue: eBay open-box deal effective price includes $5 shipping; low-feedback seller excluded",
      any(d["price"] == 34.99 and d["shipping"] == 5.0 and d["secondaryvendor"] == "Yes" for d in hd_)
      and all(d["vendor"] != "eBay" or "shadyseller" not in (d["secondaryvendorcomments"] or "") for d in hd_))
check("Hue: Qty Needed 4 plan uses the one-off eBay open-box unit once + 3 new units",
      h1.get("qtyneeded") == 4 and "1 x single @ eBay" in (h1.get("bestqtyplan") or "") and h1.get("bestqtytotal") == 184.96,
      h1.get("bestqtyplan"))
check("Plug: Shopify Product URL -> all pack variants, cheapest plan for qty 2 is the 2-pack",
      p2.get("bestqtytotal") == 26.99 and "2-pack" in (p2.get("bestqtyplan") or ""), p2.get("bestqtyplan"))
check("Plug: vendor's normal 2/4-pack pricing is not a 'deal'",
      not [d for d in deals if str(d["wishlistitem"]) == "2" and d["vendor"] == "THIRDREALITY Store"])
check("Evidence Mix recorded", "verified_direct" in (s3.get("evidencemix") or ""))

untouched = [n for n in before.sheetnames if n not in (pt.SHEET_RUN, pt.SHEET_DEALS)]
same = all([[c.value for c in row] for row in before[n].iter_rows()] == [[c.value for c in row] for row in w[n].iter_rows()]
           for n in untouched)
check(f"No other worksheet changed ({', '.join(untouched)})", same)
check("Deals Data keeps its original 8 headers in place",
      [dw.cell(1, c).value for c in range(1, 9)] ==
      ["runID", "WishlistItem", "Product", "URL", "isPrimaryVendor", "Target Price", "Price", "RightProductConfidence"])
state_dir = wb8.parent / "tracker_state"
check("State folder written (state.json + observations.jsonl)",
      (state_dir / "state.json").exists() and (state_dir / "observations.jsonl").exists())
obs = [json.loads(x) for x in (state_dir / "observations.jsonl").read_text().splitlines()]
check("Snapshot observations are never marked trusted", all(not o["trusted"] for o in obs if o["evidence"] == MARKET_SNAPSHOT) and
      any(o["trusted"] for o in obs))

n_run, n_deal, n_obs = rw.max_row, dw.max_row, len(obs)
calls_before = sum("serpapi.com/search" in u for u in fake_web.log)
pt.main(["--workbook", str(wb8), "--force", "--browser", "off", "--workers", "1", "--run-id", "test-run-1"])
w = load_workbook(wb8)
obs2 = (state_dir / "observations.jsonl").read_text().splitlines()
check("Idempotent re-run of the same run id: no duplicate Run/Deals rows or observations",
      w[pt.SHEET_RUN].max_row == n_run and w[pt.SHEET_DEALS].max_row == n_deal and len(obs2) == n_obs,
      f"{w[pt.SHEET_RUN].max_row} vs {n_run}, {w[pt.SHEET_DEALS].max_row} vs {n_deal}, {len(obs2)} vs {n_obs}")
check("Re-run inside the cache window spends no SerpApi credits", sum("serpapi.com/search" in u for u in fake_web.log) == calls_before)
dh2 = pt.header_map(w[pt.SHEET_DEALS])
ev_rerun = [w[pt.SHEET_DEALS].cell(r, dh2["priceevent"]).value for r in range(2, w[pt.SHEET_DEALS].max_row + 1)
            if w[pt.SHEET_DEALS].cell(r, dh2["runid"]).value == "test-run-1"]
check("Re-run compares against pre-run state (still 'new offer', not 'unchanged')", ev_rerun and all(e == "new offer" for e in ev_rerun), str(ev_rerun))

pt.main(["--workbook", str(wb8), "--force", "--browser", "off", "--run-id", "test-run-2", "--items", "3"])
w = load_workbook(wb8)
dw = w[pt.SHEET_DEALS]
dh2 = pt.header_map(dw)
ev2 = [(dw.cell(r, dh2["runid"]).value, dw.cell(r, dh2["priceevent"]).value) for r in range(2, dw.max_row + 1)]
check("Next run: unchanged offers are recorded but labelled 'unchanged since', not as new price events",
      [e for rid, e in ev2 if rid == "test-run-2"] and all(e.startswith("unchanged since") for rid, e in ev2 if rid == "test-run-2"),
      str(ev2))
rw = w[pt.SHEET_RUN]
rh2 = pt.header_map(rw)
last = {k: rw.cell(rw.max_row, c).value for k, c in rh2.items()}
check("Next run: prior verified price comes from the previous verified run",
      last.get("priorverifiedprice") == 219.0 and last.get("changevsprior") == 0, str(last.get("priorverifiedprice")))
check("Next run: SerpApi skipped because verified pages already price the item (credits saved)",
      "SerpApi: skipped" in (last.get("sourcenotes") or "") or "cached response" in (last.get("sourcenotes") or ""),
      last.get("sourcenotes"))

h_before = wb8.read_bytes()
pt.main(["--workbook", str(wb8), "--force", "--dry-run", "--browser", "off", "--items", "1"])
check("--dry-run leaves the workbook byte-identical", h_before == wb8.read_bytes())

# cadence: a priority-1 item checked yesterday is skipped
args = pt.parse_args([])
it1 = mk_item("X", wid=9)
it1.priority = 1
check("Cadence: priority-1 item checked yesterday is skipped",
      pt.select_items([it1], args, [{"wid": "9", "dt": datetime.utcnow() - timedelta(days=1)}], datetime.utcnow()) == [])
pt.requests.Session = _orig_session

# ---------------------------------------------------------------------------
print("\n[10] Master Sheet lists (';'), '!' exclusions, 'Only Check Primary Links'")
from pricetrack.adapters import build_query


class _Cell:                                   # openpyxl cell stand-in
    def __init__(self, v):
        self.value, self.hyperlink = v, None


check("split_multi: ';', line breaks and full-width ';' all separate entries; blanks dropped",
      pt.split_multi("a; b\nc；d;;  ;") == ["a", "b", "c", "d"], str(pt.split_multi("a; b\nc；d;;  ;")))
pos, ex = pt.split_exclusions(pt.split_multi('W50; !Lite; ! Pro Max; 50 in; 50 inch; 50"; !'))
check("'!' entries become exclusions (the phrase runs to the next ';'); a lone '!' is ignored",
      ex == ["Lite", "Pro Max"] and not any(p.startswith("!") for p in pos), f"{pos} {ex}")
check("synonym specs ('50 in', '50 inch', 50\") collapse to one phrase; a bare 'in' is read as inch",
      pos == ["W50", "50 inch"], str(pos))
check("'4 in 1' and model codes are left alone by the inch tidy-up",
      pt.tidy_spec("4 in 1 hub") == "4 in 1 hub" and pt.split_exclusions(["RAYG1US1BLK"])[0] == ["RAYG1US1BLK"])
urls = pt.extract_urls(_Cell("https://a.example.com/p; https://b.example.com/q,r ;\nhttps://c.example.com/z)"))
check("Product URLs: split on ';' / line breaks, commas INSIDE a URL are kept, trailing ')' dropped",
      urls == ["https://a.example.com/p", "https://b.example.com/q,r", "https://c.example.com/z"], str(urls))

nova = mk_item("NovaWalk W50 TrekPad with 12% auto incline", specs="W50")
nova.exclude = ["Lite", "Pro Max"]
m_ok = idn.classify(nova, "NovaWalk W50 TrekPad with 12% auto incline")
m_bad = idn.classify(nova, "NovaWalk W50 Lite TrekPad")
check("'!Lite': the plain product is High - 'Lite' is NOT required in the title", m_ok.confidence == "High", str(m_ok))
check("'!Lite': a listing containing 'Lite' is Low (excluded)", m_bad.confidence == "Low" and "excluded" in m_bad.reason, str(m_bad))
check("'!Pro Max' (multi-word phrase) excludes only the whole phrase",
      idn.excluded_phrase(nova, "NovaWalk W50 Pro Max") == "Pro Max" and idn.excluded_phrase(nova, "NovaWalk W50 Pro") is None)
check("exclusions match whole words only ('Elite' / 'Satellite' are not 'Lite')",
      idn.excluded_phrase(nova, "NovaWalk Elite W50 Satellite") is None)
sonos_x = mk_item("Sonos Ray Soundbar", "RAYG1US1BLK")
sonos_x.exclude = ["RAYG1EU1BLK"]
check("a model-code exclusion matches the listing's MPN as well as its title",
      idn.excluded_phrase(sonos_x, "Sonos Ray", mpns={"RAYG1EU1BLK"}) == "RAYG1EU1BLK")
check("Product URL page whose slug carries the excluded word is rejected",
      idn.validate_page(nova, "NovaWalk W50 TrekPad", set(), set(), slug="lite novawalk trekpad w50").confidence == "Low")
q = build_query(nova, negatives=True)
check("search query: excluded phrases are never search words, only Google minus-terms",
      q.startswith("NovaWalk W50 TrekPad") and q.endswith('-Lite -"Pro Max"') and "Lite" not in q.replace("-Lite", ""), q)
check("same-vendor query appends the merchant name before the minus-terms",
      build_query(nova, negatives=True, vendor="Best Buy") == 'NovaWalk W50 TrekPad with 12% auto incline Best Buy -Lite -"Pro Max"')
check("default query is unchanged when there are no exclusions", build_query(sonos) == "Sonos Ray Soundbar RAYG1US1BLK Unmounted")

# ---- end-to-end: links-only items --------------------------------------------------------------------------
fx10 = Path(tempfile.mkdtemp())
wb10 = fx10 / "Home Wishlist.xlsx"
shutil.copy(SRC, wb10)
w10 = load_workbook(wb10)
ms10 = w10[pt.SHEET_MASTER]
for r in range(ms10.max_row, 1, -1):
    ms10.delete_rows(r)
h10 = pt.header_map(ms10)
rows10 = [
    {"wishlistitem": 11, "product": "Zorbo Smart Lamp", "priority": 4, "quantityneeded": 1, "opentoused": "Yes",
     "targetprice": 70, "isbulkoption": "No", "onlycheckprimarylinks": "Yes",
     "producturls": "https://www.zorbo.com/products/zorbo-smart-lamp;\nhttps://www.bestbuy.com/product/zorbo-smart-lamp/J1/sku/1234567?utm_source=x"},
    {"wishlistitem": 12, "product": "Plink Wall Sconce", "priority": 4, "quantityneeded": 1, "opentoused": "No",
     "targetprice": 40, "isbulkoption": "No", "onlycheckprimarylinks": "Yes",
     "producturls": "https://www.etsy.com/listing/123/plink-wall-sconce"},
    {"wishlistitem": 13, "product": "Quill Desk Lamp", "priority": 4, "quantityneeded": 1, "opentoused": "No",
     "targetprice": 50, "isbulkoption": "No", "onlycheckprimarylinks": "Yes",
     "producturls": "https://www.quilllamps.com/products/quill-desk-lamp-old; https://www.quilllamps.com/products/quill-desk-lamp"},
    {"wishlistitem": 15, "product": "NovaWalk W50 TrekPad", "productspecifications": "W50; !Lite", "priority": 4,
     "quantityneeded": 1, "opentoused": "No", "targetprice": 300, "isbulkoption": "No", "onlycheckprimarylinks": "Yes",
     "producturls": "https://merachfit.com/products/novawalk-w50-trekpad"},
    {"wishlistitem": 16, "product": "Zeta Smart Bulb", "priority": 4, "quantityneeded": 1, "opentoused": "Yes",
     "targetprice": 20, "isbulkoption": "No", "onlycheckprimarylinks": "Yes",
     "producturls": "https://www.zeta1.com/products/zeta-smart-bulb; https://www.zeta2.com/products/zeta-smart-bulb; "
                    "https://www.zeta3.com/products/zeta-smart-bulb"},
    {"wishlistitem": 17, "product": "Orbit Desk Fan", "priority": 4, "quantityneeded": 1, "opentoused": "No",
     "targetprice": 30, "isbulkoption": "No", "onlycheckprimarylinks": "Yes"},
]
for i, row in enumerate(rows10, start=2):
    for k, v in row.items():
        col = pt.find_col(h10, k)
        if col:
            ms10.cell(i, col, v)
w10.save(wb10)


def G(title, price, source, link=None, **kw):
    d = {"title": title, "extracted_price": price, "source": source, "product_link": "https://www.google.com/shopping/product/9", **kw}
    if link:
        d["link"] = link
    return d


def serp10(url, headers):
    from urllib.parse import parse_qs, urlparse as up
    q = parse_qs(up(url).query).get("q", [""])[0]
    if "best buy" in q.lower():
        return Resp(200, {"shopping_results": [
            G("Zorbo Smart Lamp", 49.99, "Amazon.com"), G("Zorbo Smart Lamp", 52.00, "Walmart"),
            G("Case for Zorbo Smart Lamp", 9.99, "Best Buy", "https://www.bestbuy.com/site/zorbo-case/7654321.p?skuId=7654321"),
            G("Zorbo Smart Lamp", 54.99, "Best Buy", "https://www.bestbuy.com/site/zorbo-smart-lamp/1234567.p?skuId=1234567")]})
    if "etsy" in q.lower():
        return Resp(200, {"shopping_results": [
            G("Plink Wall Sconce", 24.00, "Etsy", "https://www.etsy.com/listing/123/plink-wall-sconce"),
            G("Plink Wall Sconce", 19.00, "Amazon.com")]})
    if "merachfit" in q.lower():
        return Resp(200, {"shopping_results": [
            G("NovaWalk W50 Lite TrekPad", 199.00, "MERACH"), G("NovaWalk W50 TrekPad", 279.99, "MERACH")]})
    if q == "Quill Desk Lamp":                 # general search: a row from the vendor that already priced + another seller
        return Resp(200, {"shopping_results": [G("Quill Desk Lamp", 40.00, "Quill Lamps"), G("Quill Desk Lamp", 38.00, "Amazon.com")]})
    return Resp(200, {"shopping_results": []})


zorbo_pg = ld("Zorbo Smart Lamp", "59.99", brand="Zorbo")
web10 = Web([
    ("serpapi.com/account", Resp(200, {"total_searches_left": 90})),
    ("serpapi.com/search", serp10),
    ("zorbo-smart-lamp.js", Resp(404, None, "nf")),
    ("zorbo.com/products/zorbo-smart-lamp", zorbo_pg),
    ("bestbuy.com/product/", ConnectionError("reset by peer")),
    ("bestbuy.com/site/zorbo-smart-lamp", ld("Zorbo Smart Lamp", "54.99", brand="Zorbo")),
    ("etsy.com/listing/123", Resp(403, None, "blocked")),
    ("quill-desk-lamp-old", Resp(404, None, "nf")),
    ("quilllamps.com/products/quill-desk-lamp", ld("Quill Desk Lamp", "45.00", brand="Quill")),
    ("novawalk-w50-trekpad", ld("NovaWalk W50 Lite TrekPad", "199.00", brand="Merach")),
    ("zeta-smart-bulb.js", Resp(404, None, "nf")),
    ("zeta1.com", ld("Zeta Smart Bulb", "12.00", brand="Zeta")),
    ("zeta2.com", ld("Zeta Smart Bulb", "13.00", brand="Zeta")),
    ("zeta3.com", ld("Zeta Smart Bulb", "14.00", brand="Zeta")),
])
pt.requests.Session = lambda: web10
os.environ.update({"SERPAPI_KEY": "dummy"})
os.environ.pop("BESTBUY_API_KEY", None)
rc10 = pt.main(["--workbook", str(wb10), "--force", "--browser", "off", "--workers", "1", "--run-id", "links-1",
                "--items", "11,12,13,15,16"])
pt.requests.Session = _orig_session
check("links-only run: exit code 0", rc10 == 0)
w10 = load_workbook(wb10)
rw10 = w10[pt.SHEET_RUN]
rh10 = pt.header_map(rw10)
r10 = {str(rw10.cell(r, rh10["wishlistitem"]).value): {k: rw10.cell(r, c).value for k, c in rh10.items()}
       for r in range(2, rw10.max_row + 1) if rw10.cell(r, rh10["runid"]).value == "links-1"}
for k, v in r10.items():
    print(f"     item {k}: lowest={v.get('lowestpricefound')} mix={v.get('evidencemix')} | {v.get('retrievaloutcomes')}")
    print(f"              status: {(v.get('producturlstatus') or '')[:260]}")
check("5 Run Data rows written for the links-only items", sorted(r10) == ["11", "12", "13", "15", "16"], str(sorted(r10)))
qs10 = [u for u in web10.log if "serpapi.com/search" in u]
from urllib.parse import parse_qs, urlparse as _up
queries = [parse_qs(_up(u).query)["q"][0] for u in qs10]
check("<3 working vendor links: a GENERAL search per item + a same-vendor search only for vendors with no working link",
      len(queries) == 7 and any(x.endswith("Best Buy") for x in queries) and any(x.endswith("etsy.com") for x in queries)
      and any(x.endswith("merachfit.com -Lite") for x in queries) and "Quill Desk Lamp" in queries
      and not any("quill" in x.lower() and "quilllamps" in x.lower() for x in queries), str(queries))
check("3 vendors with a working link: NO search of any kind (no SerpApi query for that item)",
      not any("zeta" in x.lower() for x in queries))
check("no sitemap / retailer search / search-engine / eBay request was made for links-only items",
      not [u for u in web10.log if any(s in u for s in ("sitemap", "robots.txt", "duckduckgo", "bing.com", "api.ebay.com",
                                                          "api.bestbuy.com", "search?", "/s/", "search_results"))],
      str([u for u in web10.log if "serpapi" not in u][:12]))
a11 = r10.get("11", {})
check("Best Buy link failed -> Best Buy's own listing found on Google, fetched and VERIFIED on the merchant page",
      "replacement Best Buy listing verified on the merchant page" in (a11.get("producturlstatus") or "")
      and "verified_discovered" in (a11.get("evidencemix") or "") and "verified_direct" in (a11.get("evidencemix") or ""),
      f"{a11.get('producturlstatus')} | {a11.get('evidencemix')}")
check("only Best Buy's listing replaces the dead Best Buy link (Amazon $49.99 / Walmart $52 / the case ignored)",
      a11.get("lowestpricefound") == 54.99 and a11.get("listingssearched") == 2, f"{a11.get('lowestpricefound')} {a11.get('listingssearched')}")
check("Source Notes tell you which link to update", "update the Master Sheet link" in (a11.get("sourcenotes") or "") and
      "Only Check Primary Links" in (a11.get("sourcenotes") or ""), a11.get("sourcenotes"))
check("skipped vendor-list search is recorded in Retrieval Outcomes (not silently dropped)", "skipped" in (a11.get("retrievaloutcomes") or ""))
check("Run Data flags: failed primary link, empty expanded search, search mode",
      (a11.get("primarylinkfailed") or "").startswith("Yes: bestbuy.com [timeout_network]")
      and (a11.get("expandedsearchnoresults") or "") == "Yes: Google Shopping (general search)"
      and "expanded SerpApi search (1 of 3 vendor links work)" in (a11.get("searchmode") or ""),
      f"{a11.get('primarylinkfailed')} | {a11.get('expandedsearchnoresults')} | {a11.get('searchmode')}")
a16 = r10.get("16", {})
check("3 working vendor links: flags read 'No' / 'Not run', mode says no search, nothing skipped silently",
      a16.get("primarylinkfailed") == "No" and (a16.get("expandedsearchnoresults") or "").startswith("Not run")
      and "no search" in (a16.get("searchmode") or "") and a16.get("listingssearched") == 3,
      f"{a16.get('primarylinkfailed')} | {a16.get('expandedsearchnoresults')} | {a16.get('searchmode')} | {a16.get('listingssearched')}")
a12 = r10.get("12", {})
check("Etsy (a Secondary-list vendor) is the item's own vendor: blocked page -> Etsy's Google row, labelled NOT verified",
      "market_snapshot, not verified" in (a12.get("producturlstatus") or "") and "market_snapshot" in (a12.get("evidencemix") or "")
      and "verified_" not in (a12.get("evidencemix") or "") and a12.get("lowestpricefound") == 24.0,
      f"{a12.get('producturlstatus')} | {a12.get('evidencemix')} | {a12.get('lowestpricefound')}")
a13 = r10.get("13", {})
check("vendor with another working Product URL: no fallback, no credit spent",
      "FAILED" in (a13.get("producturlstatus") or "") and "fallback" not in (a13.get("producturlstatus") or "")
      and a13.get("lowestpricefound") == 38.0, f"{a13.get('producturlstatus')} {a13.get('lowestpricefound')}")
check("general search runs alongside, but rows from a vendor with a working link are dropped (Amazon $38 kept, Quill $40 ignored)",
      "1 Google row(s) from vendors with a working Product URL ignored" in (a13.get("sourcenotes") or "")
      and (a13.get("expandedsearchnoresults") or "") == "No", f"{a13.get('sourcenotes')} | {a13.get('expandedsearchnoresults')}")
a15 = r10.get("15", {})
check("'!Lite': a Product URL page that is now the Lite model counts as a failed link, and the Lite Google row is not used",
      "identity_mismatch" in (a15.get("producturlstatus") or "") and "excluded by '!Lite'" in (a15.get("producturlstatus") or "")
      and a15.get("lowestpricefound") == 279.99, f"{a15.get('producturlstatus')} {a15.get('lowestpricefound')}")
check("'!Lite' went to Google as a minus-term, never as a search word",
      any("merachfit" in x and x.endswith("-Lite") and "Lite" not in x.replace("-Lite", "") for x in queries), str(queries))
c_before = len(qs10)
os.environ["PRICE_TRACKER_NO_DOTENV"] = "1"
pt.requests.Session = lambda: web10
pt.main(["--workbook", str(wb10), "--force", "--browser", "off", "--workers", "1", "--run-id", "links-2",
         "--items", "11,12,13,15,16"])
check("a re-run inside the cache window spends no further SerpApi credits",
      len([u for u in web10.log if "serpapi.com/search" in u]) == c_before)
w10b = load_workbook(wb10)
rw10b = w10b[pt.SHEET_RUN]
rh10b = pt.header_map(rw10b)
r10b = {str(rw10b.cell(r, rh10b["wishlistitem"]).value): {k: rw10b.cell(r, c).value for k, c in rh10b.items()}
        for r in range(2, rw10b.max_row + 1) if rw10b.cell(r, rh10b["runid"]).value == "links-2"}
check("Run Data flags are read back: the same failed link on the next run reads '(run 2 in a row)'",
      "run 2 in a row" in (r10b["11"].get("primarylinkfailed") or "") and "run 2 in a row" in (r10b["11"].get("expandedsearchnoresults") or "")
      and r10b["16"].get("primarylinkfailed") == "No", f"{r10b['11'].get('primarylinkfailed')} | {r10b['11'].get('expandedsearchnoresults')}")
# Only Check Primary Links = Yes with NO Product URLs -> normal search, recorded in Run Data
pt.requests.Session = lambda: web10
pt.main(["--workbook", str(wb10), "--force", "--browser", "off", "--workers", "1", "--run-id", "links-3", "--items", "17"])
pt.requests.Session = _orig_session
w10c = load_workbook(wb10)
rw10c = w10c[pt.SHEET_RUN]
rh10c = pt.header_map(rw10c)
r17 = [{k: rw10c.cell(r, c).value for k, c in rh10c.items()} for r in range(2, rw10c.max_row + 1)
       if rw10c.cell(r, rh10c["runid"]).value == "links-3"]
check("'Yes' with no Product URLs falls back to the normal search and Run Data says so",
      len(r17) == 1 and "no Product URLs are listed" in (r17[0].get("searchmode") or "")
      and "normal search used" in (r17[0].get("sourcenotes") or "") and r17[0].get("primarylinkfailed") == "No links listed",
      str(r17))

# ---------------------------------------------------------------------------
print("\n[9] No keys + no network: must not crash, must still log rows")
tmp2 = Path(tempfile.mkdtemp()) / "Home Wishlist.xlsx"
shutil.copy(FIXTURE, tmp2)
env = {k: v for k, v in os.environ.items() if k not in ("SERPAPI_KEY", "EBAY_CLIENT_ID", "EBAY_CLIENT_SECRET", "BESTBUY_API_KEY")}
proc = subprocess.run([sys.executable, "price_tracker.py", "--workbook", str(tmp2), "--force", "--no-direct", "--browser", "off"],
                      env=env, capture_output=True, text=True, timeout=300)
check("exit code 0 with no keys", proc.returncode == 0, proc.stderr[-400:])
w2 = load_workbook(tmp2)
rh = pt.header_map(w2[pt.SHEET_RUN])
n_legacy = pt.next_empty_row(load_workbook(FIXTURE)[pt.SHEET_RUN]) - 1     # (blank formatted rows are not data)
outs = [str(w2[pt.SHEET_RUN].cell(r, rh["retrievaloutcomes"]).value) for r in range(n_legacy + 1, w2[pt.SHEET_RUN].max_row + 1)]
check("Rows written; missing keys reported as api_unavailable (not a generic N/A)",
      len(outs) == 3 and all("api_unavailable" in o for o in outs), str(outs))
check("eBay is optional: no keys -> run still completes, eBay noted as optional",
      "optional" in proc.stdout.lower())

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
